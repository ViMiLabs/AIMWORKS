from __future__ import annotations

"""Generate and import the focused independent H2KG/DECODE application audit."""

import csv
import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .utils import dump_json, ensure_dir, load_json, write_text


SCHEMA_VERSION = "h2kg-independent-audit/v1"
BASELINE_VERSION = "2.3.0"
EXPECTED_COUNTS = {
    "itc_interface_field": 44,
    "public_decode_mapping": 35,
    "diameter_radius_occurrence": 7,
    "decode_local_concept": 24,
    "semantic_edge_projection": 24,
}
DECISIONS = (
    "accept_as_stated",
    "accept_with_correction",
    "reject",
    "insufficient_evidence",
)
CONFIDENCE_VALUES = ("high", "medium", "low")
DIAMETER_INTERPRETATIONS = (
    "diameter",
    "radius",
    "diameter_requires_explicit_radius_derivation",
    "insufficient_evidence",
)
REVIEW_FIELDS = (
    "decision",
    "corrected_role",
    "corrected_target_or_predicate",
    "corrected_direction_or_interpretation",
    "confidence",
    "rationale",
    "supporting_evidence",
)
EVIDENCE_FIELDS = (
    "review_item_id",
    "review_stratum",
    "item_id",
    "item_label",
    "current_role",
    "current_target",
    "current_predicate",
    "current_direction",
    "source_context",
    "evidence",
)
SHEET_NAMES = {
    "itc_interface_field": "ITC fields",
    "public_decode_mapping": "Mapping decisions",
    "diameter_radius_occurrence": "Diameter-radius cases",
    "decode_local_concept": "DECODE-local sample",
    "semantic_edge_projection": "Edge projections",
}


def build_independent_review_package(
    output_root: str | Path,
    itc_registry_path: str | Path,
    package_dir: str | Path,
    baseline_commit: str = "",
) -> dict[str, Any]:
    """Create two blank, evidence-identical reviewer workbooks and a local protocol."""
    output_root = Path(output_root)
    package_dir = ensure_dir(Path(package_dir))
    required = {
        "graph": output_root / "decode" / "decode_federated_graph.json",
        "concepts": output_root / "decode" / "decode_concept_registry.json",
        "itc": Path(itc_registry_path),
    }
    missing = [str(path) for path in required.values() if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Independent-audit inputs are missing: {missing}")

    graph = load_json(required["graph"])
    concepts = load_json(required["concepts"])
    itc = load_json(required["itc"])
    items = _build_review_items(graph, concepts, itc)
    counts = Counter(str(row["review_stratum"]) for row in items)
    if dict(counts) != EXPECTED_COUNTS:
        raise ValueError(f"Review-package coverage mismatch: expected {EXPECTED_COUNTS}, found {dict(counts)}")
    if len(items) != 134 or len({row["review_item_id"] for row in items}) != 134:
        raise ValueError("The independent audit must contain exactly 134 unique items.")

    source_fingerprints = {
        key: _file_sha256(path) for key, path in required.items()
    }
    manifest = {
        "schema_version": SCHEMA_VERSION,
        "baseline_version": BASELINE_VERSION,
        "baseline_commit": baseline_commit,
        "created_on": date.today().isoformat(),
        "source_fingerprints": source_fingerprints,
        "expected_counts": EXPECTED_COUNTS,
        "items": items,
    }
    manifest["manifest_fingerprint"] = _manifest_fingerprint(manifest)
    manifest_path = dump_json(package_dir / "independent_audit_manifest.json", manifest)

    workbooks = []
    for reviewer_code in ("R1", "R2"):
        path = package_dir / f"H2KG_DECODE_Independent_Audit_v2.3.0_{reviewer_code}.xlsx"
        _write_review_workbook(path, manifest, reviewer_code)
        workbooks.append(path)

    protocol_path = write_text(package_dir / "review_protocol.md", _review_protocol(manifest))
    email_path = write_text(package_dir / "review_invitation_email.md", _invitation_email())
    summary = {
        "status": "review_pending",
        "schema_version": SCHEMA_VERSION,
        "baseline_version": BASELINE_VERSION,
        "baseline_commit": baseline_commit,
        "review_item_count_per_reviewer": len(items),
        "stratum_counts": dict(counts),
        "manifest_fingerprint": manifest["manifest_fingerprint"],
        "manifest": str(manifest_path),
        "reviewer_workbooks": [str(path) for path in workbooks],
        "protocol": str(protocol_path),
        "invitation_email": str(email_path),
        "publication_gate": "blocked_until_two_valid_reviews_and_any_required_adjudication_are_imported",
    }
    dump_json(package_dir / "review_package_summary.json", summary)
    return summary


def import_independent_reviews(
    manifest_path: str | Path,
    reviewer_one_path: str | Path,
    reviewer_two_path: str | Path,
    public_output_dir: str | Path,
    private_output_dir: str | Path,
    adjudication_path: str | Path | None = None,
) -> dict[str, Any]:
    """Validate two authentic reviews and emit a sanitized registry or adjudication package."""
    manifest = load_json(Path(manifest_path))
    _validate_manifest(manifest)
    submissions = [
        _read_review_workbook(Path(reviewer_one_path), manifest),
        _read_review_workbook(Path(reviewer_two_path), manifest),
    ]
    if submissions[0]["identity_fingerprint"] == submissions[1]["identity_fingerprint"]:
        raise ValueError("The two review workbooks identify the same reviewer.")

    by_reviewer = [
        {row["review_item_id"]: row for row in submission["decisions"]}
        for submission in submissions
    ]
    disagreements = []
    agreement_count = 0
    for item in manifest["items"]:
        item_id = item["review_item_id"]
        first = by_reviewer[0][item_id]
        second = by_reviewer[1][item_id]
        if _decision_signature(first) == _decision_signature(second):
            agreement_count += 1
        else:
            disagreements.append({"item": item, "reviewer_one": first, "reviewer_two": second})

    public_output_dir = ensure_dir(Path(public_output_dir))
    private_output_dir = ensure_dir(Path(private_output_dir))
    private_summary = {
        "schema_version": SCHEMA_VERSION,
        "created_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "reviewers": [submission["identity"] for submission in submissions],
        "source_workbooks": [submission["workbook"] for submission in submissions],
        "source_workbook_sha256": [submission["workbook_sha256"] for submission in submissions],
        "manifest_fingerprint": manifest["manifest_fingerprint"],
    }
    dump_json(private_output_dir / "reviewer_identity_registry.private.json", private_summary)

    adjudication = None
    if disagreements:
        if adjudication_path is None:
            adjudication_workbook = private_output_dir / "H2KG_DECODE_Independent_Audit_v2.3.0_Adjudication.xlsx"
            _write_adjudication_workbook(adjudication_workbook, manifest, disagreements)
            result = {
                "status": "adjudication_required",
                "review_item_count": len(manifest["items"]),
                "initial_agreement_count": agreement_count,
                "disagreement_count": len(disagreements),
                "cohens_kappa": _cohens_kappa(submissions[0]["decisions"], submissions[1]["decisions"]),
                "adjudication_workbook": str(adjudication_workbook),
                "publication_gate": "blocked",
            }
            dump_json(public_output_dir / "independent_application_audit_status.json", result)
            return result
        adjudication = _read_adjudication_workbook(
            Path(adjudication_path), manifest, disagreements, submissions
        )
        private_summary["adjudicator"] = adjudication["identity"]
        private_summary["adjudication_workbook"] = adjudication["workbook"]
        private_summary["adjudication_workbook_sha256"] = adjudication["workbook_sha256"]
        dump_json(private_output_dir / "reviewer_identity_registry.private.json", private_summary)

    adjudicated_by_id = {
        row["review_item_id"]: row for row in (adjudication or {}).get("decisions", [])
    }
    public_rows = []
    for item in manifest["items"]:
        item_id = item["review_item_id"]
        first = by_reviewer[0][item_id]
        second = by_reviewer[1][item_id]
        agreed = _decision_signature(first) == _decision_signature(second)
        final = first if agreed else adjudicated_by_id[item_id]
        adjudicated = adjudicated_by_id.get(item_id, {})
        public_rows.append(
            {
                **item,
                "reviewer_1_decision": first["decision"],
                "reviewer_1_confidence": first["confidence"],
                "reviewer_1_rationale": first["rationale"],
                "reviewer_1_supporting_evidence": first["supporting_evidence"],
                "reviewer_2_decision": second["decision"],
                "reviewer_2_confidence": second["confidence"],
                "reviewer_2_rationale": second["rationale"],
                "reviewer_2_supporting_evidence": second["supporting_evidence"],
                "initial_agreement": agreed,
                "adjudication_required": not agreed,
                "adjudication_decision": adjudicated.get("decision", ""),
                "adjudication_confidence": adjudicated.get("confidence", ""),
                "adjudication_rationale": adjudicated.get("rationale", ""),
                "adjudication_supporting_evidence": adjudicated.get("supporting_evidence", ""),
                "final_decision": final["decision"],
                "final_corrected_role": final.get("corrected_role", ""),
                "final_corrected_target_or_predicate": final.get("corrected_target_or_predicate", ""),
                "final_corrected_direction_or_interpretation": final.get("corrected_direction_or_interpretation", ""),
                "final_rationale": final["rationale"],
                "final_disposition": _final_disposition(final["decision"]),
            }
        )

    outcome_counts = Counter(str(row["final_disposition"]) for row in public_rows)
    stratum_summary = []
    for stratum in EXPECTED_COUNTS:
        subset = [row for row in public_rows if row["review_stratum"] == stratum]
        outcomes = Counter(str(row["final_disposition"]) for row in subset)
        stratum_summary.append(
            {
                "review_stratum": stratum,
                "item_count": len(subset),
                "accepted_count": outcomes["accepted"],
                "accepted_after_correction_count": outcomes["accepted_after_correction"],
                "rejected_count": outcomes["rejected"],
                "unresolved_count": outcomes["unresolved"],
                "initial_disagreement_count": sum(not row["initial_agreement"] for row in subset),
                "adjudicated_count": sum(row["adjudication_required"] for row in subset),
            }
        )
    summary = {
        "status": "complete" if outcome_counts["unresolved"] == 0 else "complete_with_unresolved_items",
        "schema_version": SCHEMA_VERSION,
        "baseline_version": manifest["baseline_version"],
        "review_item_count": len(public_rows),
        "reviewer_count": 2 + (1 if adjudication else 0),
        "reviewer_expertise_categories": [
            submission["identity"]["expertise_category"] for submission in submissions
        ] + ([adjudication["identity"]["expertise_category"]] if adjudication else []),
        "raw_agreement_count": agreement_count,
        "raw_agreement_percent": round(100 * agreement_count / len(public_rows), 2),
        "cohens_kappa": _cohens_kappa(submissions[0]["decisions"], submissions[1]["decisions"]),
        "initial_disagreement_count": len(disagreements),
        "adjudicated_count": len(disagreements),
        "outcome_counts": dict(outcome_counts),
        "strata": stratum_summary,
        "manifest_fingerprint": manifest["manifest_fingerprint"],
        "reviewer_workbook_sha256": [submission["workbook_sha256"] for submission in submissions],
        "adjudication_workbook_sha256": adjudication["workbook_sha256"] if adjudication else "",
        "publication_gate": "passed" if outcome_counts["unresolved"] == 0 else "blocked_unresolved_items",
    }
    _write_csv(public_output_dir / "independent_application_audit.csv", public_rows)
    dump_json(
        public_output_dir / "independent_application_audit.json",
        {"summary": summary, "assessments": public_rows},
    )
    write_text(public_output_dir / "independent_application_audit.md", _audit_markdown(summary))
    dump_json(public_output_dir / "independent_application_audit_status.json", summary)
    return summary


def _build_review_items(
    graph: dict[str, Any], concepts: dict[str, Any], itc: dict[str, Any]
) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    source = itc.get("source", {})
    for field in sorted(itc.get("fields", []), key=lambda row: str(row.get("io_id", ""))):
        role = str(field.get("semantic_role", ""))
        direction = str(field.get("direction", ""))
        if direction == "input" and role == "Parameter":
            predicate_path = "hasParameter / ofProperty"
        elif direction == "input":
            predicate_path = "hasInputData / hasPart / ofProperty"
        elif role == "Data":
            predicate_path = "hasOutputData / hasPart"
        else:
            predicate_path = "hasOutputData / hasPart / ofProperty"
        items.append(
            _item(
                "itc_interface_field",
                str(field["io_id"]),
                str(field["preferred_label"]),
                role,
                "https://w3id.org/h2kg/hydrogen-ontology#" + str(field.get("h2kg_local_name", "")),
                predicate_path,
                direction,
                f"ITC interface page {field.get('page', '')}; symbol {field.get('notation', '')}; unit {field.get('reported_unit', '')}",
                f"{field.get('definition', '')} Source DOI: {source.get('scientific_source', '') or source.get('doi', '')}",
            )
        )

    public_concepts = sorted(
        (row for row in concepts.get("concepts", []) if row.get("mapping_state") == "approved_h2kg"),
        key=lambda row: str(row.get("concept_id", "")),
    )
    public_mapping_records: list[dict[str, Any]] = list(public_concepts)
    for concept in concepts.get("concepts", []):
        if concept.get("mapping_state") != "mixed_occurrence_mapping":
            continue
        matching_nodes = [
            node
            for node in graph.get("nodes", [])
            if _norm(node.get("label", "")) == _norm(concept.get("source_label", ""))
            and _norm(node.get("native_category", "")) == _norm(concept.get("native_category", ""))
            and node.get("mapping_state") == "approved_h2kg"
        ]
        for node in matching_nodes:
            public_mapping_records.append(
                {
                    **concept,
                    "concept_id": f"{concept['concept_id']}|occurrence={node['id']}",
                    "mapping_state": "approved_h2kg",
                    "anchor_iri": node.get("anchor_id", ""),
                    "semantic_role": node.get("semantic_role", ""),
                    "workflow_ids": [node.get("workflow_id", "")],
                    "rationale": node.get("mapping_evidence", ""),
                }
            )
    public_mapping_records = sorted(public_mapping_records, key=lambda row: str(row.get("concept_id", "")))
    if len(public_mapping_records) == 34:
        membrane_candidates = [
            row
            for row in concepts.get("concepts", [])
            if _norm(row.get("source_label", "")) == "membrane"
            and row.get("mapping_state") != "approved_h2kg"
        ]
        if len(membrane_candidates) != 1:
            raise ValueError(
                "Expected one DECODE-local generic Membrane concept for the withheld mapping control, "
                f"found {len(membrane_candidates)}"
            )
        membrane = membrane_candidates[0]
        public_mapping_records.append(
            {
                **membrane,
                "concept_id": f"{membrane['concept_id']}|withheld_candidate=ProtonExchangeMembrane",
                "anchor_iri": "https://w3id.org/h2kg/hydrogen-ontology#ProtonExchangeMembrane",
                "semantic_role": "Matter",
                "workflow_ids": membrane.get("workflow_ids", []),
                "rationale": (
                    "Withheld negative-control proposal. The generic source label 'Membrane' does not establish "
                    "that the material is specifically a proton-exchange membrane; no public mapping is asserted."
                ),
                "audit_assertion_status": "withheld_negative_control",
            }
        )
        public_mapping_records = sorted(public_mapping_records, key=lambda row: str(row.get("concept_id", "")))
    if len(public_mapping_records) != 35:
        raise ValueError(
            "Expected 34 asserted public DECODE mappings plus one withheld negative-control decision "
            f"in the v2.3.0 audit baseline, found {len(public_mapping_records)}"
        )
    for concept in public_mapping_records:
        concept_item_id = (
            f"{concept['concept_id']}|native_category={concept.get('native_category', '') or '[not supplied]'}"
        )
        items.append(
            _item(
                "public_decode_mapping",
                concept_item_id,
                str(concept.get("source_label", "")),
                str(concept.get("semantic_role", "")),
                str(concept.get("anchor_iri", "")),
                "mappedToAnchor",
                "source concept to public H2KG anchor",
                (
                    "WITHHELD NEGATIVE CONTROL; candidate is not asserted. Workflows: "
                    if concept.get("audit_assertion_status") == "withheld_negative_control"
                    else "ASSERTED MAPPING. Workflows: "
                )
                + " | ".join(concept.get("workflow_ids", [])),
                str(concept.get("rationale", "")),
            )
        )

    diameter_nodes = sorted(
        (
            row
            for row in graph.get("nodes", [])
            if _norm(row.get("label", "")) == "catalyst active particle diameter"
        ),
        key=lambda row: str(row.get("id", "")),
    )
    if len(diameter_nodes) != 7:
        raise ValueError(f"Expected seven affected diameter/radius occurrences, found {len(diameter_nodes)}")
    for node in diameter_nodes:
        items.append(
            _item(
                "diameter_radius_occurrence",
                str(node["id"]),
                str(node.get("label", "")),
                str(node.get("semantic_role", "")),
                str(node.get("anchor_id", "")),
                "source occurrence mapping",
                "source label interpretation",
                f"Workflow {node.get('workflow_id', '')}; source node {node.get('source_id', '')}; native category {node.get('native_category', '')}",
                str(node.get("description", "")) or "No additional source definition supplied.",
            )
        )

    local_concepts = [
        row for row in concepts.get("concepts", []) if row.get("mapping_state") != "approved_h2kg"
    ]
    role_allocations = {
        "Data": 3,
        "Instrument": 3,
        "Manufacturing": 2,
        "Matter": 3,
        "Measurement": 3,
        "Metadata": 1,
        "Parameter": 3,
        "Process": 3,
        "Property": 3,
    }
    sampled_local: list[dict[str, Any]] = []
    for role, count in role_allocations.items():
        candidates = [row for row in local_concepts if row.get("semantic_role") == role]
        if len(candidates) < count:
            raise ValueError(f"Insufficient DECODE-local concepts for role {role}: {len(candidates)}")
        sampled_local.extend(_stable_sample(candidates, count))
    for concept in sorted(sampled_local, key=lambda row: str(row.get("concept_id", ""))):
        items.append(
            _item(
                "decode_local_concept",
                (
                    f"{concept['concept_id']}|native_category="
                    f"{concept.get('native_category', '') or '[not supplied]'}"
                ),
                str(concept.get("source_label", "")),
                str(concept.get("semantic_role", "")),
                str(concept.get("anchor_iri", "")),
                "decode-local identity and coarse H2KG role",
                "not promoted to H2KG",
                "Workflows: " + " | ".join(concept.get("workflow_ids", [])),
                str(concept.get("rationale", "")),
            )
        )

    workflows = {str(row["id"]): row for row in graph.get("workflows", [])}
    edges_by_outcome: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
    for edge in graph.get("edge_registry", []):
        edges_by_outcome[str(edge.get("projection_outcome", ""))].append(edge)
    selected_edges: list[dict[str, Any]] = []
    for outcome in ("h2kg_direct", "h2kg_reified", "prov_derivation", "decode_structural_only"):
        candidates = edges_by_outcome[outcome]
        buckets: defaultdict[str, list[dict[str, Any]]] = defaultdict(list)
        for edge in candidates:
            family = str(workflows.get(str(edge.get("workflow_id", "")), {}).get("method_family", "other"))
            buckets[_workflow_bucket(family)].append(edge)
        allocation = {"manufacturing": 2, "measurement": 2, "modeling_analysis": 2}
        chosen: list[dict[str, Any]] = []
        for bucket in allocation:
            chosen.extend(_stable_sample(buckets.get(bucket, []), min(allocation[bucket], len(buckets.get(bucket, [])))))
        if len(chosen) < 6:
            selected_ids = {str(row.get("source_dependency_id", "")) for row in chosen}
            remainder = [row for row in candidates if str(row.get("source_dependency_id", "")) not in selected_ids]
            chosen.extend(_stable_sample(remainder, 6 - len(chosen)))
        if len(chosen) != 6:
            raise ValueError(f"Could not sample six edges for {outcome}")
        selected_edges.extend(chosen)
    for edge in sorted(selected_edges, key=lambda row: str(row.get("source_dependency_id", ""))):
        items.append(
            _item(
                "semantic_edge_projection",
                str(edge.get("source_dependency_id", "")),
                f"{edge.get('source_label', '')} -> {edge.get('target_label', '')}",
                f"{edge.get('source_semantic_role', '')} -> {edge.get('target_semantic_role', '')}",
                f"{edge.get('source_anchor_iri', '')} | {edge.get('target_anchor_iri', '')}",
                str(edge.get("semantic_predicates", "")) or str(edge.get("projection_outcome", "")),
                str(edge.get("semantic_direction", "")),
                f"Workflow {edge.get('workflow_id', '')}; outcome {edge.get('projection_outcome', '')}; pattern {edge.get('projection_pattern', '')}",
                str(edge.get("rationale", "")),
            )
        )
    return sorted(items, key=lambda row: (list(EXPECTED_COUNTS).index(row["review_stratum"]), row["item_id"]))


def _item(
    stratum: str,
    item_id: str,
    label: str,
    role: str,
    target: str,
    predicate: str,
    direction: str,
    context: str,
    evidence: str,
) -> dict[str, Any]:
    row = {
        "review_item_id": f"{stratum}::{item_id}",
        "review_stratum": stratum,
        "item_id": item_id,
        "item_label": label,
        "current_role": role,
        "current_target": target,
        "current_predicate": predicate,
        "current_direction": direction,
        "source_context": context,
        "evidence": evidence,
    }
    row["evidence_hash"] = _evidence_hash(row)
    return row


def _write_review_workbook(path: Path, manifest: dict[str, Any], reviewer_code: str) -> None:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instruction_rows = [
        ("H2KG/DECODE independent audit", f"Reviewer package {reviewer_code}"),
        ("Scope", "Assess all 134 rows. This is an application audit, not approval to add DECODE-local terms to H2KG."),
        ("Mapping controls", "The 35 mapping-decision rows comprise 34 asserted mappings and one clearly marked withheld negative control. Assess the evidence rather than assuming every candidate should be accepted."),
        ("Decision", "Choose accept_as_stated, accept_with_correction, reject, or insufficient_evidence."),
        ("Corrections", "For accept_with_correction, enter the corrected role, target/predicate, or direction/interpretation as applicable."),
        ("Evidence", "Enter confidence, a concise scientific rationale, and supporting evidence for every row."),
        ("Diameter/radius", "For every diameter-radius row, select one explicit interpretation in the direction/interpretation column."),
        ("H2KG Explore", "https://vimilabs.github.io/AIMWORKS/pages/explore.html"),
        ("DECODE Explorer", "https://vimilabs.github.io/AIMWORKS/pages/decode-workflows.html"),
        ("Release evidence", "https://vimilabs.github.io/AIMWORKS/pages/release.html"),
        ("Independence", "Complete the Reviewer declaration sheet. Do not consult the other reviewer."),
        ("Integrity", "Do not edit blue evidence cells. Their hashes are validated when the workbook is imported."),
        ("Return", "Return this same .xlsx file. Expected effort is approximately three to four hours."),
    ]
    for row_index, row in enumerate(instruction_rows, start=1):
        instructions.append(row)
        if str(row[1]).startswith("https://"):
            link_cell = instructions.cell(row=row_index, column=2)
            link_cell.hyperlink = row[1]
            link_cell.style = "Hyperlink"
    instructions.column_dimensions["A"].width = 24
    instructions.column_dimensions["B"].width = 110
    for row in instructions.iter_rows():
        row[0].font = Font(bold=True, color="17324D")
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    declaration = workbook.create_sheet("Reviewer declaration")
    declaration.append(["Field", "Response"])
    declaration_rows = [
        ("Reviewer name", ""),
        ("Affiliation", ""),
        ("Expertise category", ""),
        ("Independent of evaluated mappings/projections?", ""),
        ("Created any evaluated mappings/projections?", ""),
        ("Review completion date (YYYY-MM-DD)", ""),
        ("Consent to anonymized decision publication?", ""),
        ("Consent to acknowledgement by name?", ""),
    ]
    for row in declaration_rows:
        declaration.append(row)
    declaration.column_dimensions["A"].width = 48
    declaration.column_dimensions["B"].width = 65
    yes_no = DataValidation(type="list", formula1='"yes,no"', allow_blank=False)
    declaration.add_data_validation(yes_no)
    yes_no.add("B5:B6")
    yes_no.add("B8:B9")
    _style_review_sheet(declaration, unlocked_columns={2})

    for stratum, sheet_name in SHEET_NAMES.items():
        sheet = workbook.create_sheet(sheet_name)
        rows = [row for row in manifest["items"] if row["review_stratum"] == stratum]
        headers = [*EVIDENCE_FIELDS, "evidence_hash", *REVIEW_FIELDS]
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        decision_column = headers.index("decision") + 1
        confidence_column = headers.index("confidence") + 1
        interpretation_column = headers.index("corrected_direction_or_interpretation") + 1
        decision_validation = DataValidation(type="list", formula1='"' + ",".join(DECISIONS) + '"', allow_blank=False)
        confidence_validation = DataValidation(type="list", formula1='"' + ",".join(CONFIDENCE_VALUES) + '"', allow_blank=False)
        sheet.add_data_validation(decision_validation)
        sheet.add_data_validation(confidence_validation)
        decision_validation.add(f"{get_column_letter(decision_column)}2:{get_column_letter(decision_column)}{sheet.max_row}")
        confidence_validation.add(f"{get_column_letter(confidence_column)}2:{get_column_letter(confidence_column)}{sheet.max_row}")
        if stratum == "diameter_radius_occurrence":
            interpretation_validation = DataValidation(
                type="list", formula1='"' + ",".join(DIAMETER_INTERPRETATIONS) + '"', allow_blank=False
            )
            sheet.add_data_validation(interpretation_validation)
            interpretation_validation.add(
                f"{get_column_letter(interpretation_column)}2:{get_column_letter(interpretation_column)}{sheet.max_row}"
            )
        unlocked = set(range(len(EVIDENCE_FIELDS) + 2, len(headers) + 1))
        _style_review_sheet(sheet, unlocked_columns=unlocked)

    manifest_sheet = workbook.create_sheet("Manifest")
    manifest_sheet.append(["manifest_fingerprint", manifest["manifest_fingerprint"]])
    manifest_sheet.append(["schema_version", manifest["schema_version"]])
    manifest_sheet.append(["baseline_version", manifest["baseline_version"]])
    manifest_sheet.append(["reviewer_code", reviewer_code])
    manifest_sheet.append(["item_count", len(manifest["items"])])
    manifest_sheet.sheet_state = "veryHidden"
    workbook.security.lockStructure = True
    ensure_dir(path.parent)
    workbook.save(path)


def _style_review_sheet(sheet: Any, unlocked_columns: set[int]) -> None:
    header_fill = PatternFill("solid", fgColor="17324D")
    evidence_fill = PatternFill("solid", fgColor="EAF3F2")
    response_fill = PatternFill("solid", fgColor="FFF1D6")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.column in unlocked_columns:
                cell.protection = Protection(locked=False)
                cell.fill = response_fill
            else:
                cell.protection = Protection(locked=True)
                cell.fill = evidence_fill
    for column in range(1, sheet.max_column + 1):
        width = 18
        header = str(sheet.cell(1, column).value or "")
        if header in {"item_label", "source_context", "evidence", "rationale", "supporting_evidence"}:
            width = 48
        elif header in {"item_id", "current_target", "current_predicate", "corrected_target_or_predicate"}:
            width = 38
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.protection.sheet = True
    sheet.protection.password = "h2kg-audit"
    sheet.protection.autoFilter = False
    sheet.protection.sort = False


def _read_review_workbook(path: Path, manifest: dict[str, Any]) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, data_only=False)
    _validate_workbook_manifest(workbook, manifest)
    declaration = _read_declaration(workbook)
    decisions = _read_decisions(workbook, manifest)
    identity_fingerprint = hashlib.sha256(
        f"{_norm(declaration['reviewer_name'])}|{_norm(declaration['affiliation'])}".encode("utf-8")
    ).hexdigest()
    return {
        "identity": declaration,
        "identity_fingerprint": identity_fingerprint,
        "decisions": decisions,
        "workbook": str(path),
        "workbook_sha256": _file_sha256(path),
    }


def _validate_workbook_manifest(workbook: Any, manifest: dict[str, Any]) -> None:
    if "Manifest" not in workbook.sheetnames:
        raise ValueError("Review workbook has no integrity manifest.")
    values = {str(row[0].value): str(row[1].value) for row in workbook["Manifest"].iter_rows() if row[0].value}
    if values.get("manifest_fingerprint") != manifest["manifest_fingerprint"]:
        raise ValueError("Review workbook does not match the supplied evidence manifest.")
    if int(values.get("item_count", "0")) != 134:
        raise ValueError("Review workbook item count is invalid.")


def _read_declaration(workbook: Any) -> dict[str, str]:
    if "Reviewer declaration" not in workbook.sheetnames:
        raise ValueError("Reviewer declaration sheet is missing.")
    values = {
        str(row[0].value).strip(): str(row[1].value or "").strip()
        for row in workbook["Reviewer declaration"].iter_rows(min_row=2)
        if row[0].value
    }
    mapped = {
        "reviewer_name": values.get("Reviewer name", ""),
        "affiliation": values.get("Affiliation", ""),
        "expertise_category": values.get("Expertise category", ""),
        "independent": values.get("Independent of evaluated mappings/projections?", "").casefold(),
        "mapping_author": values.get("Created any evaluated mappings/projections?", "").casefold(),
        "completion_date": values.get("Review completion date (YYYY-MM-DD)", ""),
        "anonymized_publication_consent": values.get("Consent to anonymized decision publication?", "").casefold(),
        "acknowledgement_consent": values.get("Consent to acknowledgement by name?", "").casefold(),
    }
    if not mapped["reviewer_name"] or not mapped["affiliation"] or not mapped["expertise_category"]:
        raise ValueError("Reviewer name, affiliation, and expertise category are required.")
    if mapped["independent"] != "yes" or mapped["mapping_author"] != "no":
        raise ValueError("Reviewer independence declaration is not eligible for this audit.")
    if mapped["anonymized_publication_consent"] != "yes":
        raise ValueError("Consent to anonymized decision publication is required.")
    try:
        datetime.strptime(mapped["completion_date"], "%Y-%m-%d")
    except ValueError as error:
        raise ValueError("Review completion date must use YYYY-MM-DD.") from error
    if mapped["acknowledgement_consent"] not in {"yes", "no"}:
        raise ValueError("Acknowledgement consent must be yes or no.")
    return mapped


def _read_decisions(workbook: Any, manifest: dict[str, Any]) -> list[dict[str, Any]]:
    expected = {row["review_item_id"]: row for row in manifest["items"]}
    rows: list[dict[str, Any]] = []
    for stratum, sheet_name in SHEET_NAMES.items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Review sheet is missing: {sheet_name}")
        sheet = workbook[sheet_name]
        headers = [str(cell.value or "").strip() for cell in sheet[1]]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}
            item_id = str(row.get("review_item_id", ""))
            if item_id not in expected or expected[item_id]["review_stratum"] != stratum:
                raise ValueError(f"Unexpected review item: {item_id}")
            supplied_evidence = {field: row.get(field, "") for field in EVIDENCE_FIELDS}
            if _evidence_hash(supplied_evidence) != expected[item_id]["evidence_hash"]:
                raise ValueError(f"Evidence was modified for review item {item_id}")
            if str(row.get("evidence_hash", "")) != expected[item_id]["evidence_hash"]:
                raise ValueError(f"Evidence hash was modified for review item {item_id}")
            review = {field: str(row.get(field, "") or "").strip() for field in REVIEW_FIELDS}
            if review["decision"] not in DECISIONS:
                raise ValueError(f"Missing or invalid decision for {item_id}")
            if review["confidence"] not in CONFIDENCE_VALUES:
                raise ValueError(f"Missing or invalid confidence for {item_id}")
            if not review["rationale"] or not review["supporting_evidence"]:
                raise ValueError(f"Rationale and supporting evidence are required for {item_id}")
            if review["decision"] == "accept_with_correction" and not any(
                review[field] for field in (
                    "corrected_role",
                    "corrected_target_or_predicate",
                    "corrected_direction_or_interpretation",
                )
            ):
                raise ValueError(f"A correction is required for {item_id}")
            if stratum == "diameter_radius_occurrence" and review["corrected_direction_or_interpretation"] not in DIAMETER_INTERPRETATIONS:
                raise ValueError(f"An explicit diameter/radius interpretation is required for {item_id}")
            rows.append({"review_item_id": item_id, **review})
    if len(rows) != 134 or len({row["review_item_id"] for row in rows}) != 134:
        raise ValueError("A completed review must contain exactly 134 unique decisions.")
    return rows


def _write_adjudication_workbook(
    path: Path,
    manifest: dict[str, Any],
    disagreements: list[dict[str, Any]],
) -> None:
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["Purpose", "Independently adjudicate only the rows on which Reviewer 1 and Reviewer 2 disagreed."])
    instructions.append(["Requirement", "Complete the declaration and every decision row. Reviewer identities remain local."])
    declaration = workbook.create_sheet("Reviewer declaration")
    declaration.append(["Field", "Response"])
    for row in (
        ("Reviewer name", ""),
        ("Affiliation", ""),
        ("Expertise category", ""),
        ("Independent of evaluated mappings/projections?", ""),
        ("Created any evaluated mappings/projections?", ""),
        ("Review completion date (YYYY-MM-DD)", ""),
        ("Consent to anonymized decision publication?", ""),
        ("Consent to acknowledgement by name?", ""),
    ):
        declaration.append(row)
    _style_review_sheet(declaration, unlocked_columns={2})
    sheet = workbook.create_sheet("Disagreements")
    headers = [
        *EVIDENCE_FIELDS,
        "evidence_hash",
        "reviewer_1_decision",
        "reviewer_1_rationale",
        "reviewer_2_decision",
        "reviewer_2_rationale",
        *REVIEW_FIELDS,
    ]
    sheet.append(headers)
    for row in disagreements:
        item = row["item"]
        first = row["reviewer_one"]
        second = row["reviewer_two"]
        sheet.append(
            [
                *(item.get(field, "") for field in EVIDENCE_FIELDS),
                item["evidence_hash"],
                first["decision"],
                first["rationale"],
                second["decision"],
                second["rationale"],
                *("" for _field in REVIEW_FIELDS),
            ]
        )
    first_review_column = headers.index(REVIEW_FIELDS[0]) + 1
    _style_review_sheet(sheet, unlocked_columns=set(range(first_review_column, len(headers) + 1)))
    manifest_sheet = workbook.create_sheet("Manifest")
    manifest_sheet.append(["manifest_fingerprint", manifest["manifest_fingerprint"]])
    manifest_sheet.append(["schema_version", manifest["schema_version"]])
    manifest_sheet.append(["item_count", len(disagreements)])
    manifest_sheet.sheet_state = "veryHidden"
    workbook.security.lockStructure = True
    ensure_dir(path.parent)
    workbook.save(path)


def _read_adjudication_workbook(
    path: Path,
    manifest: dict[str, Any],
    disagreements: list[dict[str, Any]],
    submissions: list[dict[str, Any]],
) -> dict[str, Any]:
    if not path.exists():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, data_only=False)
    if "Manifest" not in workbook.sheetnames or "Disagreements" not in workbook.sheetnames:
        raise ValueError("Adjudication workbook is missing a required sheet.")
    values = {str(row[0].value): str(row[1].value) for row in workbook["Manifest"].iter_rows() if row[0].value}
    if values.get("manifest_fingerprint") != manifest["manifest_fingerprint"]:
        raise ValueError("Adjudication workbook does not match the audit manifest.")
    if values.get("schema_version") != manifest["schema_version"]:
        raise ValueError("Adjudication workbook schema does not match the audit manifest.")
    if int(values.get("item_count", "0")) != len(disagreements):
        raise ValueError("Adjudication workbook item count does not match the disagreements.")
    identity = _read_declaration(workbook)
    identity_fingerprint = hashlib.sha256(
        f"{_norm(identity['reviewer_name'])}|{_norm(identity['affiliation'])}".encode("utf-8")
    ).hexdigest()
    if identity_fingerprint in {submission["identity_fingerprint"] for submission in submissions}:
        raise ValueError("The adjudicator must be independent of both initial reviewers.")
    expected = {row["item"]["review_item_id"]: row for row in disagreements}
    sheet = workbook["Disagreements"]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    decisions = []
    for values_row in sheet.iter_rows(min_row=2, values_only=True):
        row = {headers[index]: values_row[index] if index < len(values_row) else "" for index in range(len(headers))}
        item_id = str(row.get("review_item_id", ""))
        if item_id not in expected:
            raise ValueError(f"Unexpected adjudication item: {item_id}")
        expected_row = expected[item_id]
        item = expected_row["item"]
        supplied = {field: row.get(field, "") for field in EVIDENCE_FIELDS}
        if _evidence_hash(supplied) != item["evidence_hash"]:
            raise ValueError(f"Adjudication evidence was modified for {item_id}")
        if str(row.get("evidence_hash", "")) != item["evidence_hash"]:
            raise ValueError(f"Adjudication evidence hash was modified for {item_id}")
        for reviewer_key, prefix in (("reviewer_one", "reviewer_1"), ("reviewer_two", "reviewer_2")):
            prior = expected_row[reviewer_key]
            if str(row.get(f"{prefix}_decision", "") or "").strip() != prior["decision"]:
                raise ValueError(f"Initial reviewer decision was modified for {item_id}")
            if str(row.get(f"{prefix}_rationale", "") or "").strip() != prior["rationale"]:
                raise ValueError(f"Initial reviewer rationale was modified for {item_id}")
        review = {field: str(row.get(field, "") or "").strip() for field in REVIEW_FIELDS}
        if review["decision"] not in DECISIONS or review["confidence"] not in CONFIDENCE_VALUES:
            raise ValueError(f"Incomplete adjudication decision for {item_id}")
        if not review["rationale"] or not review["supporting_evidence"]:
            raise ValueError(f"Adjudication rationale and evidence are required for {item_id}")
        if review["decision"] == "accept_with_correction" and not any(
            review[field] for field in (
                "corrected_role",
                "corrected_target_or_predicate",
                "corrected_direction_or_interpretation",
            )
        ):
            raise ValueError(f"An adjudication correction is required for {item_id}")
        if item["review_stratum"] == "diameter_radius_occurrence" and review["corrected_direction_or_interpretation"] not in DIAMETER_INTERPRETATIONS:
            raise ValueError(f"Adjudication must resolve diameter/radius semantics for {item_id}")
        decisions.append({"review_item_id": item_id, **review})
    if len(decisions) != len(expected) or {row["review_item_id"] for row in decisions} != set(expected):
        raise ValueError("Adjudication workbook does not cover every disagreement.")
    return {
        "identity": identity,
        "identity_fingerprint": identity_fingerprint,
        "decisions": decisions,
        "workbook": str(path),
        "workbook_sha256": _file_sha256(path),
    }


def _validate_manifest(manifest: dict[str, Any]) -> None:
    if manifest.get("schema_version") != SCHEMA_VERSION:
        raise ValueError("Unsupported independent-audit manifest schema.")
    if manifest.get("manifest_fingerprint") != _manifest_fingerprint(manifest):
        raise ValueError("Independent-audit manifest fingerprint is invalid.")
    counts = Counter(str(row.get("review_stratum", "")) for row in manifest.get("items", []))
    if dict(counts) != EXPECTED_COUNTS:
        raise ValueError(f"Independent-audit manifest coverage is invalid: {dict(counts)}")
    for row in manifest["items"]:
        if row.get("evidence_hash") != _evidence_hash(row):
            raise ValueError(f"Independent-audit evidence hash is invalid: {row.get('review_item_id', '')}")


def _manifest_fingerprint(manifest: dict[str, Any]) -> str:
    payload = {
        key: value
        for key, value in manifest.items()
        if key not in {"manifest_fingerprint", "created_on"}
    }
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest().upper()


def _evidence_hash(row: dict[str, Any]) -> str:
    payload = {field: str(row.get(field, "") or "") for field in EVIDENCE_FIELDS}
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest().upper()


def _decision_signature(row: dict[str, Any]) -> tuple[str, ...]:
    return tuple(str(row.get(field, "")).strip() for field in REVIEW_FIELDS[:4])


def _final_disposition(decision: str) -> str:
    return {
        "accept_as_stated": "accepted",
        "accept_with_correction": "accepted_after_correction",
        "reject": "rejected",
        "insufficient_evidence": "unresolved",
    }[decision]


def _cohens_kappa(first: list[dict[str, Any]], second: list[dict[str, Any]]) -> float | None:
    first_by_id = {row["review_item_id"]: row["decision"] for row in first}
    second_by_id = {row["review_item_id"]: row["decision"] for row in second}
    item_ids = sorted(set(first_by_id) & set(second_by_id))
    if not item_ids:
        return None
    observed = sum(first_by_id[item_id] == second_by_id[item_id] for item_id in item_ids) / len(item_ids)
    first_counts = Counter(first_by_id[item_id] for item_id in item_ids)
    second_counts = Counter(second_by_id[item_id] for item_id in item_ids)
    expected = sum(
        (first_counts[category] / len(item_ids)) * (second_counts[category] / len(item_ids))
        for category in DECISIONS
    )
    if expected == 1:
        return 1.0 if observed == 1 else None
    return round((observed - expected) / (1 - expected), 4)


def _workflow_bucket(family: str) -> str:
    if family == "manufacturing_assembly":
        return "manufacturing"
    if family in {
        "electrochemical_measurement",
        "electrochemical_testing",
        "imaging_characterization",
        "microscopy",
        "spectroscopy_scattering",
        "tomography",
    }:
        return "measurement"
    return "modeling_analysis"


def _stable_sample(rows: list[dict[str, Any]], limit: int) -> list[dict[str, Any]]:
    return sorted(rows, key=lambda row: hashlib.sha256(_stable_id(row).encode("utf-8")).hexdigest())[:limit]


def _stable_id(row: dict[str, Any]) -> str:
    for key in ("source_dependency_id", "concept_id", "id", "io_id", "review_item_id"):
        if row.get(key):
            return str(row[key])
    return json.dumps(row, sort_keys=True, ensure_ascii=False)


def _file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def _write_csv(path: Path, rows: list[dict[str, Any]]) -> Path:
    ensure_dir(path.parent)
    headers = list(dict.fromkeys(key for row in rows for key in row)) if rows else ["status"]
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: _cell(row.get(key, "")) for key in headers})
    return path


def _cell(value: Any) -> Any:
    if isinstance(value, (list, tuple, set)):
        return " | ".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def _norm(value: Any) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(value).casefold()).split())


def _review_protocol(manifest: dict[str, Any]) -> str:
    lines = [
        "# H2KG/DECODE independent application audit protocol",
        "",
        f"Baseline: H2KG {manifest['baseline_version']} ({manifest.get('baseline_commit') or 'commit recorded at package generation'})",
        f"Manifest fingerprint: `{manifest['manifest_fingerprint']}`",
        "",
        "Two eligible experts review the same 134 items independently. The workbooks contain identical, protected evidence and no prefilled decisions. A third expert receives only disagreement rows.",
        "",
        "| Stratum | Items |",
        "|---|---:|",
    ]
    for stratum, count in EXPECTED_COUNTS.items():
        lines.append(f"| {stratum} | {count} |")
    lines.extend(
        [
            "| **Total** | **134** |",
            "",
            "A completed submission requires all decisions, confidence ratings, rationales, supporting evidence, and declarations. Evidence edits, duplicate identities, mapping authorship, and unresolved disagreement prevent import.",
            "",
            "The mapping-decision stratum contains 34 asserted mappings and one withheld negative control (generic Membrane to ProtonExchangeMembrane). The withheld candidate is not part of the released public mapping count.",
            "",
            "DECODE-local acceptance confirms only project-local identity and semantic role. It does not create an H2KG class or W3ID identifier.",
        ]
    )
    return "\n".join(lines) + "\n"


def _invitation_email() -> str:
    return """Subject: Independent review of H2KG/DECODE semantic application items

Dear [Reviewer name],

I would like to invite you to complete a focused independent assessment of 134 prespecified H2KG/DECODE application items. The attached workbook is the complete review instrument; you do not need to read the manuscript or inspect the full 5,486-row automated consistency inventory.

Please complete the declaration and every orange response field. For each item, select a decision, record confidence, provide a concise rationale and supporting evidence, and enter a correction where needed. The evidence cells are protected so the returned workbook can be verified against its source fingerprint. Please work independently and do not consult the other reviewer.

The mapping-decision sheet contains 34 asserted mappings and one explicitly marked withheld negative control. Please judge each row from its evidence; the control is included to test whether the protocol distinguishes a generic source term from an unsupported narrower target.

The expected effort is approximately three to four hours, and a return within ten days would be appreciated. Acceptance of a DECODE-local item concerns its project-level identity and role only; it does not approve that concept for addition to H2KG.

Kind regards,
[Name]
"""


def _audit_markdown(summary: dict[str, Any]) -> str:
    lines = [
        "# Independent application audit",
        "",
        f"Status: **{summary['status']}**",
        "",
        f"Two independent reviewers assessed {summary['review_item_count']} prespecified items. Raw agreement was {summary['raw_agreement_percent']:.2f}% and Cohen's kappa was {summary['cohens_kappa']}.",
        "",
        "| Stratum | Items | Accepted | Corrected | Rejected | Unresolved | Initial disagreements |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for row in summary["strata"]:
        lines.append(
            f"| {row['review_stratum']} | {row['item_count']} | {row['accepted_count']} | "
            f"{row['accepted_after_correction_count']} | {row['rejected_count']} | {row['unresolved_count']} | "
            f"{row['initial_disagreement_count']} |"
        )
    return "\n".join(lines) + "\n"
