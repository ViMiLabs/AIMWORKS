from __future__ import annotations

import json
import shutil
from pathlib import Path

from openpyxl import load_workbook

from aimworks_ontology_release.application_evaluation import build_application_evaluation


def test_application_evaluation_is_reproducible_and_traceable(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    output = tmp_path / "output"
    for relative in (
        "decode/decode_federated_graph.json",
        "decode/decode_concept_registry.json",
        "ontology/taxonomy_registry.json",
        "ontology/external_alignment_registry.json",
        "ontology/h2kg-v2-complete.ttl",
        "benchmarks/tib_collections/ontology_benchmark.csv",
    ):
        source = root / "output" / relative
        target = output / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)

    result = build_application_evaluation(
        output,
        root / "input" / "itc_curated_interface.json",
        root / "input" / "tib_collection_candidate_cache.json",
        root / "config" / "comparator_evaluation.json",
    )
    assert result["status"] == "passed"
    assert result["coverage"]["canonical_source_concepts"] == 573
    assert result["coverage"]["public_h2kg_outcomes"] == 29
    assert result["coverage"]["decode_local_outcomes"] == 544
    assert result["itc_decode_crosswalk"]["itc_field_count"] == 44
    assert result["itc_decode_crosswalk"]["fields_with_concept_wide_source_alignment"] == 6
    assert result["itc_decode_crosswalk"]["fields_without_concept_wide_alignment"] == 38
    assert result["itc_decode_crosswalk"]["fields_with_occurrence_specific_source_alignment"] == 5
    assert result["itc_decode_crosswalk"]["fields_without_source_occurrence"] == 33
    assert result["itc_decode_crosswalk"]["source_occurrences_covered_by_concept_wide_fields"] == 14
    assert result["itc_decode_crosswalk"]["source_occurrences_covered_by_occurrence_specific_fields"] == 5
    assert result["itc_decode_crosswalk"]["source_occurrences_covered_total"] == 19

    crosswalk = json.loads(
        (output / "evaluation" / "itc_decode_population_crosswalk.json").read_text(encoding="utf-8")
    )
    radius = next(
        row for row in crosswalk["rows"]
        if row["io_id"] == "IO_MAT_CatalystActiveParticleDiameter"
    )
    assert radius["population_layer"] == "source_occurrence_specific_alignment"
    assert radius["source_occurrence_count"] == 1
    assert "iet-ink-to-composition-model::n6" in radius["occurrence_ids"]
    assert result["independent_semantic_audit_status"] == "review_pending"
    assert result["independent_judgments_completed"] == 0
    assert result["independent_application_audit"]["review_item_count"] == 134
    strata = {
        row["check_stratum"]: row["item_count"]
        for row in result["automated_consistency"]["strata"]
    }
    assert strata["itc_field"] == 44
    assert strata["decode_concept"] == 573
    assert strata["decode_edge"] == 4463
    assert strata["h2kg_taxonomy_sample"] == 150
    assert result["automated_consistency"]["check_count"] == 5486
    assert result["retrieval"]["rq1_target_input_concept_count"] == 24
    assert result["retrieval"]["rq1_common_population_output_occurrence_count"] == 917
    assert result["retrieval"]["rq1_semantic_result_count"] == 3
    assert result["retrieval"]["rq1_semantic_workflow_count"] == 1
    assert result["retrieval"]["rq1_exact_label_result_count"] == 15
    assert result["retrieval"]["rq1_exact_label_workflow_count"] == 9
    assert result["retrieval"]["rq2_scalar_candidate_interpretation"].endswith("not established false positives")
    assert result["retrieval"]["independent_precision_recall_claimed"] is False
    assert result["comparator_evaluation"]["resource_count"] == 4
    assert result["comparator_evaluation"]["task_count"] == 5
    assert result["comparator_evaluation"]["assessment_count"] == 20
    assert result["comparator_evaluation"]["aggregate_score_reported"] is False

    workbook = load_workbook(output / "evaluation" / "automated_consistency_inventory.xlsx")
    assert "ITC fields" in workbook.sheetnames
    assert workbook["ITC fields"].max_row == 45
    assert workbook["DECODE concepts"].max_row == 574
    assert workbook["DECODE edges"].max_row == 4464
    assert not (output / "evaluation" / "semantic_assessment_registry.json").exists()
    retrieval = (output / "evaluation" / "retrieval_evaluation.json").read_text(encoding="utf-8")
    assert "evaluation-reviewer" not in retrieval
    assert "independently labelled" not in retrieval
    comparator = load_workbook(output / "evaluation" / "focused_comparator_matrix.xlsx")
    assert "Comparator matrix" in comparator.sheetnames
