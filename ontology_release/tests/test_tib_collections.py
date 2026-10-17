from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from aimworks_ontology_release.tib_collections import (
    OUTCOMES,
    _assess_candidate,
    _external_role_evidence,
    _write_csv,
    apply_harmonization_registry,
    build_harmonization_registry,
    build_collection_benchmark,
    summarize_decisions,
)


H2KG = "https://w3id.org/h2kg/hydrogen-ontology#"


def test_write_csv_removes_trailing_whitespace_from_embedded_lines(tmp_path: Path) -> None:
    target = tmp_path / "evidence.csv"
    _write_csv(target, [{"definition": "first line  \nsecond line \t"}], ["definition"])

    content = target.read_text(encoding="utf-8")
    assert "first line  \n" not in content
    assert "second line \t" not in content
    assert "first line\nsecond line" in content


class FakeResponse:
    def __init__(self, payload: dict):
        self.payload = payload

    def raise_for_status(self) -> None:
        return None

    def json(self) -> dict:
        return self.payload


class FakeSession:
    def get(self, url: str, params: dict, timeout: int):
        if url.endswith("/ontologies"):
            return FakeResponse(
                {
                    "_embedded": {
                        "ontologies": [
                            {
                                "ontologyId": "chmo",
                                "version": "1",
                                "numberOfTerms": 100,
                                "config": {"title": "CHMO", "collection": ["NFDI4CAT", "NFDI4ING"]},
                            },
                            {
                                "ontologyId": "oeo",
                                "version": "2",
                                "numberOfTerms": 200,
                                "config": {"title": "Open Energy Ontology", "collection": ["NFDI4ENERGY"]},
                            },
                        ]
                    }
                }
            )
        if "/search" in url:
            query = params["q"].casefold()
            if "transmission electron microscopy" in query or query == "tem":
                return FakeResponse(
                    {
                        "response": {
                            "docs": [
                                {
                                    "ontology_name": "chmo",
                                    "iri": "urn:chmo:tem",
                                    "label": "transmission electron microscopy",
                                    "description": ["A microscopy method using transmitted electrons."],
                                    "type": "class",
                                }
                            ]
                        }
                    }
                )
            if "afm scan speed" in query:
                return FakeResponse(
                    {
                        "response": {
                            "docs": [
                                {
                                    "ontology_name": "chmo",
                                    "iri": "urn:chmo:afm",
                                    "label": "atomic force microscopy",
                                    "description": ["A probe microscopy method."],
                                    "type": "class",
                                }
                            ]
                        }
                    }
                )
            return FakeResponse({"response": {"docs": []}})
        if "/classes/" in url and "urn%253Achmo%253Atem" in url:
            return FakeResponse(
                {
                    "iri": "urn:chmo:tem",
                    "label": ["transmission electron microscopy"],
                    "definition": [{"value": "A microscopy method using transmitted electrons."}],
                    "definedBy": ["chmo"],
                    "appearsIn": ["chmo", "oeo"],
                    "isDefiningOntology": True,
                    "imported": False,
                    "hierarchicalAncestor": ["urn:obi:assay"],
                    "linkedEntities": {"urn:obi:assay": {"label": ["assay"]}},
                }
            )
        if "/classes/" in url and "urn%253Achmo%253Aafm" in url:
            return FakeResponse(
                {
                    "iri": "urn:chmo:afm",
                    "label": ["atomic force microscopy"],
                    "definedBy": ["chmo"],
                    "appearsIn": ["chmo"],
                    "isDefiningOntology": True,
                    "imported": False,
                    "hierarchicalAncestor": ["urn:obi:assay"],
                    "linkedEntities": {"urn:obi:assay": {"label": ["assay"]}},
                }
            )
        raise AssertionError(f"Unexpected request: {url} {params}")


def _registry(path: Path) -> Path:
    terms = [
        {
            "iri": f"{H2KG}TransmissionElectronMicroscopyImaging",
            "local_name": "TransmissionElectronMicroscopyImaging",
            "label": "Transmission Electron Microscopy Imaging",
            "alt_labels": ["TEM imaging"],
            "definition": "A measurement that acquires images using transmitted electrons.",
            "source_roles": ["Measurement"],
            "corrected_semantic_role": "Measurement",
            "disposition": "active_local_class",
            "parent_label": "Electron microscopy measurement",
            "top_branch": "Process",
            "ancestor_path": [f"{H2KG}Process", f"{H2KG}Measurement"],
            "module": "core",
        },
        {
            "iri": f"{H2KG}AFMScanSpeed",
            "local_name": "AFMScanSpeed",
            "label": "AFM Scan Speed",
            "alt_labels": [],
            "definition": "A scan-speed parameter used for AFM acquisition.",
            "source_roles": ["Parameter"],
            "corrected_semantic_role": "Characteristic",
            "disposition": "active_local_class",
            "parent_label": "Acquisition characteristic",
            "top_branch": "Characteristic",
            "ancestor_path": [f"{H2KG}Characteristic"],
            "module": "core",
        },
        {
            "iri": f"{H2KG}IonomerToCarbonRatio",
            "local_name": "IonomerToCarbonRatio",
            "label": "Ionomer-to-carbon ratio",
            "alt_labels": ["I/C ratio"],
            "definition": "A composition ratio.",
            "source_roles": ["Parameter"],
            "corrected_semantic_role": "Characteristic",
            "disposition": "active_local_class",
            "parent_label": "Composition characteristic",
            "top_branch": "Characteristic",
            "ancestor_path": [f"{H2KG}Characteristic"],
            "module": "core",
        },
        {
            "iri": f"{H2KG}DeprecatedClass",
            "local_name": "DeprecatedClass",
            "label": "Deprecated class",
            "disposition": "deprecated_compatibility_class",
        },
    ]
    path.write_text(
        json.dumps(
            {
                "schema_version": "h2kg-taxonomy-registry/v2",
                "version": "2.0.0",
                "source_sha256": "TEST",
                "terms": terms,
            }
        ),
        encoding="utf-8",
    )
    return path


def test_collection_benchmark_generates_percentages_excel_and_safe_overlay(tmp_path: Path) -> None:
    registry = _registry(tmp_path / "taxonomy_registry.json")
    summary = build_collection_benchmark(
        registry,
        tmp_path / "output",
        tmp_path / "tib_collection_snapshot.json",
        tmp_path / "tib_collection_candidates.json",
        refresh=True,
        session=FakeSession(),
    )
    target = tmp_path / "output" / "benchmarks" / "tib_collections"

    assert summary["h2kg_active_classes_assessed"] == 3
    assert (target / "h2kg_tib_collection_benchmark.xlsx").exists()
    assert (target / "h2kg_tib_manuscript_tables.xlsx").exists()
    assert (target / "ontology_summary.csv").exists()
    assert (target / "ontology_benchmark.csv").exists()
    assert (target / "manuscript_table.md").exists()
    workbook = load_workbook(target / "h2kg_tib_manuscript_tables.xlsx", read_only=False)
    assert workbook.sheetnames == [
        "Read me",
        "Collection summary",
        "Ontology benchmark",
        "H2KG branch summary",
    ]

    benchmark_summary = json.loads((target / "benchmark_summary.json").read_text(encoding="utf-8"))
    assert benchmark_summary["ontology_class_assignments"] == (
        benchmark_summary["h2kg_active_classes_assessed"] * benchmark_summary["ontology_rows"]
    )
    assert benchmark_summary["ontology_class_assignments"] == benchmark_summary["expected_ontology_class_assignments"]
    rows = benchmark_summary["collection_summary"]
    nfdi4cat = next(row for row in rows if row["collection"] == "NFDI4Cat")
    assert nfdi4cat["machine_exact_equivalent_count"] == 1
    assert nfdi4cat["machine_exact_equivalent_percent"] == 33.3333
    assert round(sum(nfdi4cat[f"{outcome}_percent"] for outcome in OUTCOMES), 2) == 100.0

    harmonization = json.loads((target / "automatic_harmonization_registry.json").read_text(encoding="utf-8"))
    assert any(row["local_iri"].endswith("TransmissionElectronMicroscopyImaging") for row in harmonization["mappings"])
    assert harmonization["label_changes"][0]["new_preferred_label"] == "Transmission electron microscopy"
    assert harmonization["owl_imports_added"] == []

    imported = (target / "imported_availability.csv").read_text(encoding="utf-8")
    assert "oeo" in imported
    decisions = (target / "candidate_decisions.csv").read_text(encoding="utf-8")
    assert "AFM Scan Speed" not in decisions or "machine_exact_equivalent" not in next(
        line for line in decisions.splitlines() if "AFM Scan Speed" in line
    )


def test_incomplete_retrieval_blocks_publication(tmp_path: Path) -> None:
    registry = _registry(tmp_path / "taxonomy_registry.json")
    snapshot = {
        "retrieved_at": "2026-10-14T00:00:00Z",
        "collections": {
            "nfdi4cat": {
                "catalogue_name": "NFDI4CAT",
                "requested_name": "NFDI4CAT",
                "ontology_ids": ["chmo"],
            }
        },
        "ontologies": [{"ontology_id": "chmo", "title": "CHMO", "number_of_terms": 1}],
    }
    cache = {
        "format": "h2kg-tib-collection-candidates/v1",
        "complete": False,
        "query_failure_count": 1,
        "detail_failure_count": 0,
        "queries": {
            "manufacturing": {
                "query": "Manufacturing",
                "status": "error",
                "error": "TIB unavailable",
            }
        },
        "details": {},
        "candidates": [],
    }
    snapshot_path = tmp_path / "snapshot.json"
    cache_path = tmp_path / "cache.json"
    snapshot_path.write_text(json.dumps(snapshot), encoding="utf-8")
    cache_path.write_text(json.dumps(cache), encoding="utf-8")
    with pytest.raises(RuntimeError, match="publication blocked"):
        build_collection_benchmark(
            registry,
            tmp_path / "output",
            snapshot_path,
            cache_path,
        )
    target = tmp_path / "output" / "benchmarks" / "tib_collections"
    assert "blocking" in (target / "retrieval_failures.csv").read_text(encoding="utf-8")


def test_harmonization_registry_is_installed_without_imports(tmp_path: Path) -> None:
    source = tmp_path / "generated.json"
    source.write_text(
        json.dumps(
            {
                "format": "h2kg-tib-auto-harmonization/v1",
                "mappings": [{"status": "machine_safe"}],
                "label_changes": [],
                "owl_imports_added": [],
            }
        ),
        encoding="utf-8",
    )
    target = tmp_path / "config" / "tib_auto_harmonization.json"
    result = apply_harmonization_registry(source, target)
    assert result["status"] == "applied"
    assert result["owl_imports_added"] == 0
    assert target.exists()


def test_harmonization_abstains_for_multiple_exact_targets_in_one_ontology() -> None:
    term = {
        "h2kg_iri": f"{H2KG}TransmissionElectronMicroscopyImaging",
        "h2kg_label": "Transmission Electron Microscopy Imaging",
        "h2kg_role": "Measurement",
    }
    base = {
        **term,
        "ontology_id": "chmo",
        "collection_memberships": "NFDI4CAT",
        "external_label": "transmission electron microscopy",
        "outcome": "machine_exact_equivalent",
        "safe_for_harmonization": True,
        "native_to_ontology": True,
        "decision_basis": "Test exact match.",
        "lexical_similarity": 100,
    }
    registry = build_harmonization_registry(
        [term],
        [
            {**base, "external_iri": "urn:chmo:tem-1"},
            {**base, "external_iri": "urn:chmo:tem-2"},
        ],
        {"retrieved_at": "2026-10-14T00:00:00Z"},
        {"version": "2.0.0", "source_sha256": "TEST"},
    )
    assert registry["mappings"] == []
    assert registry["label_changes"] == []


def test_failed_api_assessment_blocks_final_semantic_summary() -> None:
    term = {
        "h2kg_iri": f"{H2KG}Manufacturing",
        "h2kg_label": "Manufacturing",
        "h2kg_role": "Process",
        "h2kg_top_branch": "Process",
    }
    snapshot = {
        "retrieved_at": "2026-10-14T00:00:00Z",
        "collections": {
            "nfdi4ing": {
                "catalogue_name": "NFDI4ING",
                "requested_name": "NFDI4ING",
                "ontology_ids": ["emmo"],
            }
        },
        "ontologies": [{"ontology_id": "emmo", "title": "EMMO", "number_of_terms": 1}],
    }
    with pytest.raises(RuntimeError, match="Retrieval failure"):
        summarize_decisions(
            [term],
            [],
            snapshot,
            incomplete_term_iris={term["h2kg_iri"]},
        )


def test_many_h2kg_classes_to_one_external_target_are_related_not_equivalent() -> None:
    terms = [
        {
            "h2kg_iri": f"{H2KG}SEMImaging",
            "h2kg_label": "SEM Imaging",
            "h2kg_role": "Measurement",
            "h2kg_top_branch": "Process",
        },
        {
            "h2kg_iri": f"{H2KG}SEMImagingMeasurement",
            "h2kg_label": "SEM Imaging Measurement",
            "h2kg_role": "Measurement",
            "h2kg_top_branch": "Process",
        },
    ]
    decisions = [
        {
            **term,
            "ontology_id": "chmo",
            "collection_memberships": "NFDI4CAT",
            "external_iri": "urn:chmo:sem",
            "external_label": "scanning electron microscopy",
            "outcome": "machine_exact_equivalent",
            "safe_for_harmonization": True,
            "native_to_ontology": True,
            "decision_basis": "Test exact match.",
            "lexical_similarity": 100,
        }
        for term in terms
    ]
    snapshot = {
        "retrieved_at": "2026-10-14T00:00:00Z",
        "collections": {
            "nfdi4cat": {
                "catalogue_name": "NFDI4CAT",
                "requested_name": "NFDI4CAT",
                "ontology_ids": ["chmo"],
            }
        },
        "ontologies": [{"ontology_id": "chmo", "title": "CHMO", "number_of_terms": 1}],
    }
    summary = summarize_decisions(terms, decisions, snapshot)
    assert summary["collection_summary"][0]["ambiguous_abstained_count"] == 0
    assert summary["collection_summary"][0]["related_non_equivalent_count"] == 2
    registry = build_harmonization_registry(
        terms,
        decisions,
        snapshot,
        {"version": "2.0.0", "source_sha256": "TEST"},
    )
    assert registry["mappings"] == []


def _snapshot() -> dict:
    return {
        "collections": {
            "nfdi4cat": {
                "catalogue_name": "NFDI4CAT",
                "requested_name": "NFDI4CAT",
                "ontology_ids": ["chmo"],
            }
        },
        "ontologies": [{"ontology_id": "chmo", "title": "CHMO"}],
    }


def _term(local_name: str, label: str, role: str, definition: str) -> dict:
    return {
        "h2kg_iri": f"{H2KG}{local_name}",
        "h2kg_local_name": local_name,
        "h2kg_label": label,
        "h2kg_alt_labels": [],
        "h2kg_definition": definition,
        "h2kg_role": role,
        "h2kg_top_branch": "Process" if role in {"Process", "Measurement"} else role,
        "h2kg_parent_label": "",
        "h2kg_module": "core",
    }


def _candidate(label: str, definition: str, ancestors: list[str]) -> dict:
    return {
        "ontology_id": "chmo",
        "iri": f"urn:chmo:{label.replace(' ', '-')}",
        "label": label,
        "definition": definition,
        "synonyms": [],
        "type": "class",
        "native_to_ontology": True,
        "is_obsolete": False,
        "is_defining_ontology": True,
        "ancestor_labels": ancestors,
        "ancestor_iris": [f"urn:ancestor:{index}" for index, _ in enumerate(ancestors)],
    }


def test_atomic_layer_deposition_is_process_and_exact_equivalent() -> None:
    candidate = _candidate(
        "atomic layer deposition",
        "A thin-film deposition technique based on sequential surface reactions.",
        ["sample deposition", "synthesis method", "process"],
    )
    evidence = _external_role_evidence(candidate)
    assert evidence["role"] == "Process"
    decision = _assess_candidate(
        _term(
            "AtomicLayerDeposition",
            "Atomic Layer Deposition",
            "Process",
            "A coating or deposition process using sequential self-limiting surface reactions.",
        ),
        candidate,
        _snapshot(),
    )
    assert decision["final_outcome"] == "machine_exact_equivalent"


@pytest.mark.parametrize(
    ("term", "candidate", "expected"),
    [
        (
            _term("AemionIonomer", "Aemion ionomer", "Matter", "An Aemion-brand anion-exchange ionomer material."),
            _candidate("ionomer", "A polymer material containing ion-conducting groups.", ["polymer material"]),
            "h2kg_narrower",
        ),
        (
            _term("AFMTipNominalRadius", "AFM tip nominal radius", "Characteristic", "The nominal radius of an AFM probe tip."),
            _candidate("AFM tip", "A probe component used by an atomic force microscope.", ["instrument component"]),
            "related_non_equivalent",
        ),
        (
            _term("ContactAngleMeasurement", "Contact angle measurement", "Measurement", "A measurement of contact angle."),
            _candidate("contact angle", "An angle characteristic at a liquid-solid interface.", ["angle characteristic"]),
            "related_non_equivalent",
        ),
        (
            _term("AcidEtching", "Acid etching", "Process", "An etching process that uses an acid reagent."),
            _candidate("etching", "A material-removal process.", ["material processing"]),
            "h2kg_narrower",
        ),
    ],
)
def test_directional_and_related_regression_cases(term: dict, candidate: dict, expected: str) -> None:
    assert _assess_candidate(term, candidate, _snapshot())["final_outcome"] == expected


def test_exact_alias_does_not_erase_material_state_or_property_scope() -> None:
    solution = _term(
        "AceticAcidSolution",
        "Acetic acid solution",
        "Matter",
        "A diluted acetic acid solution used as a peptiser.",
    )
    solution["h2kg_alt_labels"] = ["acetic acid"]
    acid = _candidate(
        "acetic acid",
        "A monocarboxylic acid containing two carbon atoms.",
        ["monocarboxylic acid", "chemical entity"],
    )
    solution_decision = _assess_candidate(solution, acid, _snapshot())
    assert solution_decision["final_outcome"] == "related_non_equivalent"
    assert solution_decision["material_state_conflict"] is True
    assert solution_decision["safe_for_harmonization"] is False

    acid_value = _term(
        "AcidValue",
        "Acid value",
        "Characteristic",
        "A property determined by acid-base titration and reported as neutralization amount per sample mass.",
    )
    acid_value["h2kg_alt_labels"] = ["acidity"]
    acidity = _candidate(
        "acidity",
        "The tendency of a compound to donate a proton.",
        ["physico-chemical property"],
    )
    acidity_decision = _assess_candidate(acid_value, acidity, _snapshot())
    assert acidity_decision["final_outcome"] == "related_non_equivalent"
    assert acidity_decision["safe_for_harmonization"] is False


def test_broad_search_alias_does_not_create_exact_equivalence() -> None:
    average = _term(
        "AveragePoreDiameter",
        "Average pore diameter",
        "Characteristic",
        "The arithmetic or distribution-derived mean diameter of pores in a sample.",
    )
    average["h2kg_alt_labels"] = ["pore diameter"]
    pore_diameter = _candidate(
        "pore diameter",
        "The diameter of an individual pore.",
        ["diameter", "length characteristic"],
    )
    decision = _assess_candidate(average, pore_diameter, _snapshot())
    assert decision["final_outcome"] == "h2kg_narrower"
    assert decision["final_outcome"] != "machine_exact_equivalent"
    assert "alias identity is discovery evidence only" in decision["scope_evidence"]


def test_camel_case_method_label_is_not_misclassified_as_instrument() -> None:
    candidate = _candidate(
        "ScanningElectronMicroscopy",
        "A microscopy technique in which a scanning electron microscope produces surface images.",
        ["Microscopy", "CharacterisationTechnique"],
    )
    evidence = _external_role_evidence(candidate)
    assert evidence["role"] == "Measurement"
    decision = _assess_candidate(
        _term(
            "SEMInstrument",
            "SEM instrument",
            "Instrument",
            "An electron microscope used for scanning-electron-microscopy acquisition.",
        ),
        candidate,
        _snapshot(),
    )
    assert decision["final_outcome"] == "related_non_equivalent"
    assert decision["safe_for_harmonization"] is False


def test_agent_entity_is_not_exact_to_agent_role() -> None:
    candidate = _candidate(
        "agent",
        "An agent is a role of a person or organisation directing activity towards goals.",
        ["role", "realizable entity"],
    )
    candidate["direct_parent_labels"] = ["role"]
    decision = _assess_candidate(
        _term(
            "Agent",
            "Agent",
            "Agent",
            "An entity that bears responsibility for an activity or generated resource.",
        ),
        candidate,
        _snapshot(),
    )
    assert decision["final_outcome"] == "related_non_equivalent"
    assert decision["granularity_conflict"] is True
    assert decision["safe_for_harmonization"] is False


@pytest.mark.parametrize(
    ("external_label", "external_definition"),
    [
        ("concentration unit", "A unit used to report molar concentration."),
        ("reaction temperature profile", "A data profile of reaction temperature over time."),
    ],
)
def test_parameter_is_not_asserted_against_unit_or_profile(
    external_label: str,
    external_definition: str,
) -> None:
    term = _term(
        "ReactionTemperature",
        "Reaction temperature",
        "Characteristic",
        "A parameter specifying the temperature of a reaction.",
    )
    term["h2kg_source_roles"] = ["Parameter"]
    candidate = _candidate(
        external_label,
        external_definition,
        ["information entity", "entity"],
    )
    decision = _assess_candidate(term, candidate, _snapshot())
    assert decision["granularity_conflict"] is True
    assert decision["safe_for_harmonization"] is False


def test_high_confidence_directional_relations_use_h2kg_subject_convention() -> None:
    narrower_term = _term(
        "AemionIonomer",
        "Aemion ionomer",
        "Matter",
        "An Aemion-brand anion-exchange ionomer material.",
    )
    broader_term = _term(
        "Ammonia",
        "Ammonia",
        "Matter",
        "A material entity corresponding to ammonia without a phase restriction.",
    )
    narrower = _assess_candidate(
        narrower_term,
        _candidate("ionomer", "A polymer material containing ion-conducting groups.", ["polymer material"]),
        _snapshot(),
    )
    broader = _assess_candidate(
        broader_term,
        _candidate("ammonia gas", "A gas composed of ammonia molecules.", ["gas", "material"]),
        _snapshot(),
    )
    registry = build_harmonization_registry(
        [narrower_term, broader_term],
        [narrower, broader],
        {**_snapshot(), "retrieved_at": "2026-10-14T00:00:00Z"},
        {"version": "2.0.0", "source_sha256": "TEST"},
    )
    relations = {row["local_iri"]: row["relation"] for row in registry["mappings"]}
    assert relations[narrower_term["h2kg_iri"]] == "skos:broadMatch"
    assert relations[broader_term["h2kg_iri"]] == "skos:narrowMatch"
