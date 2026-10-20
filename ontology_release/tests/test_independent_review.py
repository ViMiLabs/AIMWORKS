from __future__ import annotations

import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from aimworks_ontology_release.independent_review import (
    SHEET_NAMES,
    build_independent_review_package,
    import_independent_reviews,
)


def _complete_workbook(path: Path, reviewer: str, affiliation: str) -> None:
    workbook = load_workbook(path)
    declaration = workbook["Reviewer declaration"]
    responses = [
        reviewer,
        affiliation,
        "hydrogen electrochemistry and scientific data",
        "yes",
        "no",
        "2026-10-20",
        "yes",
        "no",
    ]
    for row, value in enumerate(responses, start=2):
        declaration.cell(row=row, column=2, value=value)
    for stratum, sheet_name in SHEET_NAMES.items():
        sheet = workbook[sheet_name]
        headers = {str(cell.value): cell.column for cell in sheet[1]}
        for row in range(2, sheet.max_row + 1):
            sheet.cell(row, headers["decision"], "accept_as_stated")
            sheet.cell(row, headers["confidence"], "high")
            sheet.cell(row, headers["rationale"], "The supplied identity, role, and relation are scientifically coherent.")
            sheet.cell(row, headers["supporting_evidence"], "Assessment based on the supplied source context and ontology evidence.")
            if stratum == "diameter_radius_occurrence":
                sheet.cell(row, headers["corrected_direction_or_interpretation"], "diameter")
    workbook.save(path)


def _complete_adjudication_workbook(path: Path, reviewer: str, affiliation: str) -> None:
    workbook = load_workbook(path)
    declaration = workbook["Reviewer declaration"]
    responses = [
        reviewer,
        affiliation,
        "hydrogen electrochemistry and ontology evaluation",
        "yes",
        "no",
        "2026-10-20",
        "yes",
        "no",
    ]
    for row, value in enumerate(responses, start=2):
        declaration.cell(row=row, column=2, value=value)
    sheet = workbook["Disagreements"]
    headers = {str(cell.value): cell.column for cell in sheet[1]}
    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, headers["decision"], "accept_as_stated")
        sheet.cell(row, headers["confidence"], "high")
        sheet.cell(row, headers["rationale"], "The supplied role and mapping are supported by the source context.")
        sheet.cell(row, headers["supporting_evidence"], "Independent adjudication of the supplied evidence.")
        if sheet.cell(row, headers["review_stratum"]).value == "diameter_radius_occurrence":
            sheet.cell(row, headers["corrected_direction_or_interpretation"], "diameter")
    workbook.save(path)


def test_review_package_has_exact_scope_and_blank_decisions(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    package_dir = tmp_path / "review"
    result = build_independent_review_package(
        root / "output",
        root / "input" / "itc_curated_interface.json",
        package_dir,
        baseline_commit="test-commit",
    )
    assert result["status"] == "review_pending"
    assert result["review_item_count_per_reviewer"] == 134
    assert result["stratum_counts"] == {
        "itc_interface_field": 44,
        "public_decode_mapping": 35,
        "diameter_radius_occurrence": 7,
        "decode_local_concept": 24,
        "semantic_edge_projection": 24,
    }
    manifest = json.loads((package_dir / "independent_audit_manifest.json").read_text(encoding="utf-8"))
    assert len(manifest["items"]) == 134
    assert len({row["review_item_id"] for row in manifest["items"]}) == 134
    workbook = load_workbook(package_dir / "H2KG_DECODE_Independent_Audit_v2.3.0_R1.xlsx")
    assert workbook["ITC fields"].max_row == 45
    assert workbook["Mapping decisions"].max_row == 36
    assert workbook["Diameter-radius cases"].max_row == 8
    assert workbook["Manifest"].sheet_state == "veryHidden"
    headers = {str(cell.value): cell.column for cell in workbook["ITC fields"][1]}
    assert workbook["ITC fields"].cell(2, headers["decision"]).value is None
    assert workbook["ITC fields"].cell(2, headers["decision"]).protection.locked is False
    assert workbook["ITC fields"].cell(2, headers["evidence"]).protection.locked is True


def test_review_import_validates_identity_and_evidence(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    package_dir = tmp_path / "review"
    build_independent_review_package(
        root / "output",
        root / "input" / "itc_curated_interface.json",
        package_dir,
    )
    first = package_dir / "H2KG_DECODE_Independent_Audit_v2.3.0_R1.xlsx"
    second = package_dir / "H2KG_DECODE_Independent_Audit_v2.3.0_R2.xlsx"
    _complete_workbook(first, "Reviewer One", "Institute One")
    _complete_workbook(second, "Reviewer Two", "Institute Two")
    result = import_independent_reviews(
        package_dir / "independent_audit_manifest.json",
        first,
        second,
        tmp_path / "public",
        tmp_path / "private",
    )
    assert result["status"] == "complete"
    assert result["review_item_count"] == 134
    assert result["raw_agreement_percent"] == 100.0
    public_text = (tmp_path / "public" / "independent_application_audit.json").read_text(encoding="utf-8")
    assert "Reviewer One" not in public_text
    assert "Institute One" not in public_text
    assert "Reviewer Two" not in public_text
    private_text = (tmp_path / "private" / "reviewer_identity_registry.private.json").read_text(encoding="utf-8")
    assert "Reviewer One" in private_text

    tampered = tmp_path / "tampered.xlsx"
    shutil.copy2(first, tampered)
    workbook = load_workbook(tampered)
    workbook["ITC fields"]["J2"] = "Altered evidence"
    workbook.save(tampered)
    with pytest.raises(ValueError, match="Evidence was modified"):
        import_independent_reviews(
            package_dir / "independent_audit_manifest.json",
            tampered,
            second,
            tmp_path / "public-2",
            tmp_path / "private-2",
        )


def test_duplicate_reviewer_identity_is_rejected(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    package_dir = tmp_path / "review"
    build_independent_review_package(
        root / "output",
        root / "input" / "itc_curated_interface.json",
        package_dir,
    )
    first = package_dir / "H2KG_DECODE_Independent_Audit_v2.3.0_R1.xlsx"
    second = package_dir / "H2KG_DECODE_Independent_Audit_v2.3.0_R2.xlsx"
    _complete_workbook(first, "Same Reviewer", "Same Institute")
    _complete_workbook(second, "Same Reviewer", "Same Institute")
    with pytest.raises(ValueError, match="same reviewer"):
        import_independent_reviews(
            package_dir / "independent_audit_manifest.json",
            first,
            second,
            tmp_path / "public",
            tmp_path / "private",
        )


def test_disagreement_blocks_publication_until_independent_adjudication(tmp_path: Path) -> None:
    root = Path(__file__).resolve().parents[1]
    package_dir = tmp_path / "review"
    build_independent_review_package(
        root / "output",
        root / "input" / "itc_curated_interface.json",
        package_dir,
    )
    first = package_dir / "H2KG_DECODE_Independent_Audit_v2.3.0_R1.xlsx"
    second = package_dir / "H2KG_DECODE_Independent_Audit_v2.3.0_R2.xlsx"
    _complete_workbook(first, "Reviewer One", "Institute One")
    _complete_workbook(second, "Reviewer Two", "Institute Two")

    workbook = load_workbook(second)
    sheet = workbook["ITC fields"]
    headers = {str(cell.value): cell.column for cell in sheet[1]}
    sheet.cell(2, headers["decision"], "reject")
    sheet.cell(2, headers["rationale"], "The proposed semantic identity is not supported by the supplied evidence.")
    workbook.save(second)

    blocked = import_independent_reviews(
        package_dir / "independent_audit_manifest.json",
        first,
        second,
        tmp_path / "public",
        tmp_path / "private",
    )
    assert blocked["status"] == "adjudication_required"
    assert blocked["publication_gate"] == "blocked"
    assert blocked["disagreement_count"] == 1
    assert not (tmp_path / "public" / "independent_application_audit.json").exists()

    adjudication_path = Path(blocked["adjudication_workbook"])
    adjudication = load_workbook(adjudication_path)
    assert adjudication["Disagreements"].max_row == 2
    _complete_adjudication_workbook(adjudication_path, "Reviewer Three", "Institute Three")

    completed = import_independent_reviews(
        package_dir / "independent_audit_manifest.json",
        first,
        second,
        tmp_path / "public-final",
        tmp_path / "private-final",
        adjudication_path=adjudication_path,
    )
    assert completed["status"] == "complete"
    assert completed["reviewer_count"] == 3
    assert completed["initial_disagreement_count"] == 1
    assert completed["adjudicated_count"] == 1
    public_text = (tmp_path / "public-final" / "independent_application_audit.json").read_text(encoding="utf-8")
    assert "Reviewer Three" not in public_text
    assert "Institute Three" not in public_text
    private_text = (tmp_path / "private-final" / "reviewer_identity_registry.private.json").read_text(encoding="utf-8")
    assert "Reviewer Three" in private_text
