from __future__ import annotations

import csv
import json
import re
import shutil
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Iterable

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from rdflib import Graph, OWL, RDF, RDFS, URIRef

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - build dependency guard
    raise SystemExit("Install the manuscript extra: pip install -e ontology_release[manuscript]") from exc


HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
ONTOLOGY_RELEASE = REPO / "ontology_release"
FIGURES = HERE / "figures"
TABLES = HERE / "tables"

H2KG = "https://w3id.org/h2kg/hydrogen-ontology#"
QUDT = "http://qudt.org/schema/qudt/"

COLORS = {
    "ink": "#172B3A",
    "muted": "#627586",
    "line": "#536A79",
    "paper": "#FCFBF7",
    "matter": "#DCEFEB",
    "matter_edge": "#18796D",
    "process": "#DCEAF7",
    "process_edge": "#2B6EA6",
    "measurement": "#D9E9F8",
    "measurement_edge": "#23689B",
    "parameter": "#FFF0D5",
    "parameter_edge": "#B26D00",
    "data": "#EEE4F8",
    "data_edge": "#8151AB",
    "property": "#DCEFF3",
    "property_edge": "#337C8D",
    "metadata": "#ECEFF2",
    "metadata_edge": "#6A7680",
    "accent": "#D46A3A",
    "related": "#96A5AF",
    "unmatched": "#D9DEE2",
}


def configure_plotting() -> None:
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Aptos", "Segoe UI", "DejaVu Sans"],
            "font.size": 10,
            "axes.titlesize": 16,
            "axes.titleweight": "bold",
            "figure.facecolor": COLORS["paper"],
            "axes.facecolor": COLORS["paper"],
            "svg.fonttype": "none",
        }
    )


def draw_box(
    ax,
    x: float,
    y: float,
    width: float,
    height: float,
    title: str,
    subtitle: str = "",
    *,
    face: str = "#FFFFFF",
    edge: str = COLORS["line"],
    title_size: float = 10.5,
    subtitle_size: float = 8.5,
    dashed: bool = False,
    zorder: int = 3,
) -> FancyBboxPatch:
    patch = FancyBboxPatch(
        (x, y),
        width,
        height,
        boxstyle="round,pad=0.012,rounding_size=0.018",
        linewidth=1.5,
        edgecolor=edge,
        facecolor=face,
        linestyle="--" if dashed else "-",
        zorder=zorder,
    )
    ax.add_patch(patch)
    ax.text(
        x + width / 2,
        y + height * (0.61 if subtitle else 0.5),
        title,
        ha="center",
        va="center",
        fontsize=title_size,
        fontweight="bold",
        color=COLORS["ink"],
        zorder=zorder + 1,
        wrap=True,
    )
    if subtitle:
        ax.text(
            x + width / 2,
            y + height * 0.29,
            subtitle,
            ha="center",
            va="center",
            fontsize=subtitle_size,
            color=COLORS["muted"],
            zorder=zorder + 1,
            wrap=True,
        )
    return patch


def arrow(
    ax,
    start: tuple[float, float],
    end: tuple[float, float],
    label: str = "",
    *,
    color: str = COLORS["line"],
    dashed: bool = False,
    rad: float = 0.0,
    label_offset: tuple[float, float] = (0.0, 0.0),
    fontsize: float = 8,
) -> None:
    patch = FancyArrowPatch(
        start,
        end,
        arrowstyle="-|>",
        mutation_scale=12,
        linewidth=1.25,
        color=color,
        linestyle="--" if dashed else "-",
        connectionstyle=f"arc3,rad={rad}",
        zorder=2,
    )
    ax.add_patch(patch)
    if label:
        ax.text(
            (start[0] + end[0]) / 2 + label_offset[0],
            (start[1] + end[1]) / 2 + label_offset[1],
            label,
            ha="center",
            va="center",
            fontsize=fontsize,
            color=color,
            bbox={"boxstyle": "round,pad=0.13", "facecolor": COLORS["paper"], "edgecolor": "none", "alpha": 0.96},
            zorder=5,
        )


def panel_label(ax, label: str, x: float, y: float) -> None:
    ax.text(x, y, label, fontsize=15, fontweight="bold", color=COLORS["ink"], va="top")


def save_figure(fig, stem: str) -> None:
    FIGURES.mkdir(parents=True, exist_ok=True)
    svg_path = FIGURES / f"{stem}.svg"
    fig.savefig(svg_path, bbox_inches="tight", facecolor=fig.get_facecolor())
    svg_text = svg_path.read_text(encoding="utf-8")
    svg_path.write_text("\n".join(line.rstrip() for line in svg_text.splitlines()) + "\n", encoding="utf-8")
    fig.savefig(FIGURES / f"{stem}.png", dpi=400, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)


def figure_1_scope_architecture() -> None:
    fig, ax = plt.subplots(figsize=(12.8, 7.1))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("H₂KG scope, foundational reuse and external alignment", loc="left", pad=16)

    draw_box(ax, 0.035, 0.27, 0.23, 0.46, "Hydrogen electrochemical\nscope", "Materials · manufacturing\nmeasurements · models\ndata · properties · provenance", face="#F2F5F2", edge="#6C8574", title_size=11)
    draw_box(ax, 0.37, 0.37, 0.26, 0.27, "H₂KG 2.2", "2,221 active classes\nfive intrinsic branches\nstable w3id.org identifiers", face="#E7F0F3", edge="#176D79", title_size=18, subtitle_size=10)

    foundations = [
        (0.30, "EMMO / ECHO / HDO"),
        (0.47, "ChEBI"),
        (0.60, "QUDT"),
        (0.72, "PROV-O"),
        (0.84, "DCTERMS"),
    ]
    for x, label in foundations:
        draw_box(ax, x, 0.82, 0.115 if x > 0.3 else 0.155, 0.09, label, face="#F4F6F7", edge=COLORS["metadata_edge"], title_size=8.7)
        arrow(ax, (x + (0.077 if x == 0.30 else 0.057), 0.82), (0.5, 0.65), color="#7D8A93", rad=0.05)

    draw_box(ax, 0.72, 0.46, 0.24, 0.18, "External SKOS bridge", "Exact · H₂KG narrower\nH₂KG broader", face="#F4EEF8", edge=COLORS["data_edge"], title_size=12)
    draw_box(ax, 0.72, 0.19, 0.24, 0.16, "Frozen TIB evidence", "NFDI4Cat · NFDI4Energy\nNFDI4Ing", face="#FFF4E4", edge=COLORS["parameter_edge"], title_size=11)
    arrow(ax, (0.63, 0.505), (0.72, 0.55), "publishes", label_offset=(0, 0.03))
    arrow(ax, (0.84, 0.35), (0.84, 0.46), "gates", color=COLORS["parameter_edge"], label_offset=(0.04, 0))
    arrow(ax, (0.265, 0.50), (0.37, 0.50), "domain requirements")

    outputs = [
        (0.10, "Taxonomy", "OWL / RDF"),
        (0.285, "Profiles", "SHACL"),
        (0.47, "Explore", "human navigation"),
    ]
    for x, title, subtitle in outputs:
        draw_box(ax, x, 0.055, 0.16, 0.105, title, subtitle, face="#FFFFFF", edge=COLORS["process_edge"], title_size=10)
        arrow(ax, (0.50, 0.37), (x + 0.08, 0.16), color=COLORS["process_edge"], rad=0.06)

    ax.text(0.035, 0.94, "FOUNDATIONAL REUSE", fontsize=9, fontweight="bold", color=COLORS["muted"])
    ax.text(0.72, 0.69, "ALIGNMENT, NOT COLLECTION-WIDE IMPORT", fontsize=9, fontweight="bold", color=COLORS["muted"])
    ax.text(0.035, 0.015, "Ontology terminology remains separate from workflow occurrence records.", fontsize=9, color=COLORS["muted"])
    save_figure(fig, "figure_1_scope_architecture")


def figure_2_engineering_workflow() -> None:
    fig, ax = plt.subplots(figsize=(13.2, 6.4))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("H₂KG 2.2 engineering, benchmarking and release workflow", loc="left", pad=16)

    boxes = [
        (0.03, "1  Requirements", "Workflow scope\ncompetency questions"),
        (0.195, "2  Term registry", "Identifiers · labels\ndefinitions · provenance"),
        (0.36, "3  Taxonomy correction", "Explicit parents\nsplits · replacements"),
        (0.525, "4  Role migration", "Intrinsic identity\nSHACL workflow use"),
        (0.69, "5  TIB comparison", "Frozen evidence\nconservative bridge"),
        (0.855, "6  Validated release", "OWL · JSON-LD\nExplore · reports"),
    ]
    for i, (x, title, subtitle) in enumerate(boxes):
        face = ["#F3F1E9", "#E9F2EF", "#E8F0F7", "#FFF1DA", "#F1E8F7", "#E5F0F1"][i]
        edge = ["#7A7356", COLORS["matter_edge"], COLORS["process_edge"], COLORS["parameter_edge"], COLORS["data_edge"], COLORS["property_edge"]][i]
        draw_box(ax, x, 0.48, 0.135, 0.24, title, subtitle, face=face, edge=edge, title_size=10.2, subtitle_size=8.4)
        if i:
            arrow(ax, (x - 0.03, 0.60), (x, 0.60), color=COLORS["line"])

    checks = [
        (0.10, "Structural", "cycles · orphans\nlabels · definitions"),
        (0.29, "Logical", "ELK classification\ndomain/range"),
        (0.48, "Constraint", "SHACL positive\nand negative tests"),
        (0.67, "Serialization", "Turtle ↔ JSON-LD\nRDF isomorphism"),
        (0.86, "Competency", "9 executable\nretrieval checks"),
    ]
    for x, title, subtitle in checks:
        draw_box(ax, x - 0.065, 0.12, 0.13, 0.16, title, subtitle, face="#FFFFFF", edge=COLORS["metadata_edge"], title_size=9.5, subtitle_size=7.8)
        arrow(ax, (x, 0.28), (x, 0.47), color=COLORS["metadata_edge"], dashed=True)

    ax.text(0.03, 0.84, "AUTHORITATIVE INPUTS", fontsize=9, fontweight="bold", color=COLORS["muted"])
    ax.text(0.03, 0.34, "RELEASE GATES", fontsize=9, fontweight="bold", color=COLORS["muted"])
    ax.text(0.03, 0.035, "A failed gate blocks publication; a failed TIB request is never converted into an unmatched class.", fontsize=9, color=COLORS["muted"])
    save_figure(fig, "figure_2_engineering_workflow")


def figure_3_taxonomy_role_pattern() -> None:
    fig, ax = plt.subplots(figsize=(12.8, 8.0))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("Intrinsic taxonomy and contextual workflow roles", loc="left", pad=16)

    panel_label(ax, "a", 0.025, 0.91)
    draw_box(ax, 0.40, 0.79, 0.20, 0.10, "H₂KG active taxonomy", "2,221 classes", face="#E7F0F3", edge="#176D79", title_size=12)
    branches = [
        (0.025, "Agent", "4", COLORS["metadata"], COLORS["metadata_edge"]),
        (0.215, "Characteristic", "1,052", COLORS["property"], COLORS["property_edge"]),
        (0.405, "Information Entity", "109", COLORS["data"], COLORS["data_edge"]),
        (0.595, "Matter", "692", COLORS["matter"], COLORS["matter_edge"]),
        (0.785, "Process", "364", COLORS["process"], COLORS["process_edge"]),
    ]
    for x, title, count, face, edge in branches:
        draw_box(ax, x, 0.59, 0.165, 0.12, title, f"{count} active classes", face=face, edge=edge, title_size=10.5)
        arrow(ax, (0.50, 0.79), (x + 0.0825, 0.71), color=edge)

    ax.text(0.225, 0.545, "Property", ha="center", fontsize=9, color=COLORS["property_edge"], fontweight="bold")
    ax.text(0.445, 0.545, "Parameter · Data · DataPoint · Metadata", ha="center", fontsize=8.6, color=COLORS["data_edge"], fontweight="bold")
    ax.text(0.675, 0.545, "Instrument", ha="center", fontsize=9, color=COLORS["matter_edge"], fontweight="bold")
    ax.text(0.867, 0.545, "Manufacturing · Measurement", ha="center", fontsize=8.6, color=COLORS["process_edge"], fontweight="bold")

    panel_label(ax, "b", 0.025, 0.47)
    ax.text(0.065, 0.458, "Example: ionomer-to-carbon ratio used as a model input", fontsize=12, fontweight="bold", color=COLORS["ink"], va="top")
    draw_box(ax, 0.04, 0.20, 0.17, 0.14, "Ink-to-Composition\nmodel run", "Process occurrence", face=COLORS["process"], edge=COLORS["process_edge"])
    draw_box(ax, 0.29, 0.20, 0.16, 0.14, "I/C ratio record", "Parameter record", face=COLORS["parameter"], edge=COLORS["parameter_edge"])
    draw_box(ax, 0.53, 0.20, 0.18, 0.14, "I/C ratio occurrence", "IonomerToCarbonRatio", face=COLORS["property"], edge=COLORS["property_edge"])
    draw_box(ax, 0.79, 0.20, 0.16, 0.14, "Quantity value", "numeric value + unit", face="#FFF6E8", edge=COLORS["parameter_edge"])
    arrow(ax, (0.21, 0.27), (0.29, 0.27), "hasParameter", color=COLORS["process_edge"], label_offset=(0, 0.035))
    arrow(ax, (0.45, 0.27), (0.53, 0.27), "ofProperty", color=COLORS["parameter_edge"], label_offset=(0, 0.035))
    arrow(ax, (0.71, 0.27), (0.79, 0.27), "hasQuantityValue", color=COLORS["property_edge"], label_offset=(0, 0.035))
    draw_box(ax, 0.32, 0.035, 0.38, 0.085, "Intrinsic identity remains Characteristic → Property", "Workflow role is expressed by the occurrence path, not by reclassifying the property.", face="#FFFFFF", edge=COLORS["property_edge"], title_size=10, subtitle_size=8)
    arrow(ax, (0.62, 0.20), (0.56, 0.12), "rdf:type", color=COLORS["property_edge"], dashed=True, label_offset=(0.04, 0), fontsize=7.5)
    save_figure(fig, "figure_3_taxonomy_role_pattern")


def figure_4_tib_benchmark() -> None:
    summary = json.loads((ONTOLOGY_RELEASE / "output/benchmarks/tib_collections/benchmark_summary.json").read_text(encoding="utf-8"))
    rows = summary["collection_summary"]
    categories = [
        ("machine_exact_equivalent_percent", "Exact equivalent", "#176D79"),
        ("h2kg_narrower_percent", "H₂KG narrower", "#2B6EA6"),
        ("h2kg_broader_percent", "H₂KG broader", "#D3822A"),
        ("related_non_equivalent_percent", "Related", "#8A69A6"),
        ("unmatched_percent", "Unmatched", COLORS["unmatched"]),
    ]
    fig, ax = plt.subplots(figsize=(11.8, 5.9))
    names = [row["collection"] for row in rows]
    y = list(range(len(rows)))
    left = [0.0] * len(rows)
    for key, label, color in categories:
        values = [row[key] for row in rows]
        bars = ax.barh(y, values, left=left, height=0.55, label=label, color=color, edgecolor=COLORS["paper"], linewidth=0.8)
        for index, (bar, value) in enumerate(zip(bars, values)):
            if value >= 3.0:
                ax.text(left[index] + value / 2, bar.get_y() + bar.get_height() / 2, f"{value:.2f}%", ha="center", va="center", fontsize=8.2, color="#FFFFFF" if key != "unmatched_percent" else COLORS["ink"], fontweight="bold")
        left = [a + b for a, b in zip(left, values)]
    ax.set_yticks(y, names, fontsize=11, fontweight="bold")
    ax.invert_yaxis()
    ax.set_xlim(0, 100)
    ax.set_xlabel("Percentage of 2,221 active H₂KG classes")
    ax.set_title("Automated correspondence with TIB collection unions", loc="left", pad=16)
    ax.grid(axis="x", color="#CBD2D6", linewidth=0.7, alpha=0.7)
    ax.set_axisbelow(True)
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.legend(ncol=5, loc="lower center", bbox_to_anchor=(0.5, -0.29), frameon=False, fontsize=9)
    fig.text(0.01, 0.01, "Labels are shown for segments ≥3%; Table 3 reports all values. Exact-IRI reuse and ambiguous outcomes are zero.", fontsize=8.5, color=COLORS["muted"])
    fig.subplots_adjust(left=0.16, right=0.98, top=0.84, bottom=0.24)
    save_figure(fig, "figure_4_tib_benchmark")


def figure_5_synthetic_tem() -> None:
    fig, ax = plt.subplots(figsize=(14.0, 8.8))
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    ax.set_title("Synthetic TEM workflow represented with H₂KG", loc="left", pad=16)

    ax.text(0.99, 0.96, "SYNTHETIC DEMONSTRATION · NO EXPERIMENTAL CLAIM", ha="right", va="center", fontsize=9.5, fontweight="bold", color="#A55422", bbox={"boxstyle": "round,pad=0.35", "facecolor": "#FFF1DA", "edgecolor": "#D3822A"})

    panel_label(ax, "a", 0.02, 0.89)
    ax.text(0.055, 0.88, "Specimen preparation", fontsize=12, fontweight="bold", color=COLORS["ink"], va="top")
    draw_box(ax, 0.025, 0.70, 0.11, 0.09, "Pd/C", "Matter", face=COLORS["matter"], edge=COLORS["matter_edge"])
    draw_box(ax, 0.025, 0.56, 0.11, 0.09, "Liquid mixture", "Matter", face=COLORS["matter"], edge=COLORS["matter_edge"], title_size=9.4)
    draw_box(ax, 0.19, 0.625, 0.14, 0.12, "Mix + sonicate", "Manufacturing", face=COLORS["process"], edge=COLORS["process_edge"])
    draw_box(ax, 0.39, 0.625, 0.12, 0.12, "Dispersion", "Matter", face=COLORS["matter"], edge=COLORS["matter_edge"])
    draw_box(ax, 0.58, 0.625, 0.14, 0.12, "Deposit + dry", "Manufacturing", face=COLORS["process"], edge=COLORS["process_edge"])
    draw_box(ax, 0.58, 0.79, 0.14, 0.075, "TEM grid", "Matter input", face=COLORS["matter"], edge=COLORS["matter_edge"], title_size=9.5, subtitle_size=7.3)
    draw_box(ax, 0.80, 0.625, 0.13, 0.12, "TEM sample", "Matter", face=COLORS["matter"], edge=COLORS["matter_edge"])
    arrow(ax, (0.19, 0.70), (0.135, 0.745), "hasInputMaterial", color=COLORS["process_edge"], fontsize=6.7, label_offset=(-0.002, 0.022))
    arrow(ax, (0.19, 0.67), (0.135, 0.605), "hasInputMaterial", color=COLORS["process_edge"], fontsize=6.7, label_offset=(-0.002, -0.018))
    arrow(ax, (0.33, 0.685), (0.39, 0.685), "hasOutputMaterial", color=COLORS["process_edge"], fontsize=6.7, label_offset=(0, 0.034))
    arrow(ax, (0.58, 0.685), (0.51, 0.685), "hasInputMaterial", color=COLORS["process_edge"], fontsize=6.7, label_offset=(0, 0.034))
    arrow(ax, (0.65, 0.745), (0.65, 0.79), "hasInputMaterial", color=COLORS["process_edge"], fontsize=6.7, label_offset=(0.052, 0))
    arrow(ax, (0.72, 0.685), (0.80, 0.685), "hasOutputMaterial", color=COLORS["process_edge"], fontsize=6.7, label_offset=(0, 0.034))
    ax.text(0.965, 0.565, "Grouped preparation activities; all arrows follow RDF direction.", fontsize=7.2, color=COLORS["muted"], ha="right")

    panel_label(ax, "b", 0.02, 0.49)
    ax.text(0.055, 0.48, "Acquisition, data processing and result provenance", fontsize=12, fontweight="bold", color=COLORS["ink"], va="top")
    main_y, main_h = 0.285, 0.105
    draw_box(ax, 0.02, main_y, 0.10, main_h, "TEM sample", "Matter", face=COLORS["matter"], edge=COLORS["matter_edge"], title_size=9.3)
    draw_box(ax, 0.155, main_y, 0.145, main_h, "TEM measurement (M)", "Transmission electron\nmicroscopy imaging", face=COLORS["measurement"], edge=COLORS["measurement_edge"], title_size=9.5, subtitle_size=7.2)
    draw_box(ax, 0.34, main_y, 0.105, main_h, "Raw TEM images", "Data", face=COLORS["data"], edge=COLORS["data_edge"], title_size=9.0)
    draw_box(ax, 0.485, main_y, 0.10, main_h, "Format\nconversion", "Process", face=COLORS["process"], edge=COLORS["process_edge"], title_size=8.8)
    draw_box(ax, 0.625, main_y, 0.105, main_h, "Processed\nimages", "Data", face=COLORS["data"], edge=COLORS["data_edge"], title_size=8.8)
    draw_box(ax, 0.77, main_y, 0.095, main_h, "Particle\nanalysis (A)", "Process", face=COLORS["process"], edge=COLORS["process_edge"], title_size=8.5)
    draw_box(ax, 0.905, main_y, 0.075, main_h, "Result\ndata", "Data", face=COLORS["data"], edge=COLORS["data_edge"], title_size=8.6)

    arrow(ax, (0.155, 0.337), (0.12, 0.337), "hasInputMaterial", color=COLORS["measurement_edge"], fontsize=6.4, label_offset=(0, 0.031))
    arrow(ax, (0.30, 0.337), (0.34, 0.337), "hasOutputData", color=COLORS["measurement_edge"], fontsize=6.4, label_offset=(0, 0.031))
    arrow(ax, (0.485, 0.337), (0.445, 0.337), "hasInputData", color=COLORS["process_edge"], fontsize=6.4, label_offset=(0, 0.031))
    arrow(ax, (0.585, 0.337), (0.625, 0.337), "hasOutputData", color=COLORS["process_edge"], fontsize=6.4, label_offset=(0, 0.031))
    arrow(ax, (0.77, 0.337), (0.73, 0.337), "hasInputData", color=COLORS["process_edge"], fontsize=6.4, label_offset=(0, 0.031))
    arrow(ax, (0.865, 0.337), (0.905, 0.337), "hasOutputData", color=COLORS["process_edge"], fontsize=6.4, label_offset=(0, 0.031))

    draw_box(ax, 0.015, 0.105, 0.105, 0.085, "Metadata", "synthetic context", face=COLORS["metadata"], edge=COLORS["metadata_edge"], title_size=8.8, subtitle_size=7.0, dashed=True)
    draw_box(ax, 0.155, 0.105, 0.12, 0.085, "TEM instrument", "Instrument", face=COLORS["metadata"], edge=COLORS["metadata_edge"], title_size=8.8, subtitle_size=7.0)
    draw_box(ax, 0.31, 0.105, 0.13, 0.085, "Parameter records", "six synthetic settings", face=COLORS["parameter"], edge=COLORS["parameter_edge"], title_size=8.8, subtitle_size=7.0)
    draw_box(ax, 0.47, 0.105, 0.10, 0.085, "5 nm", "synthetic QUDT value", face="#FFF5E5", edge=COLORS["parameter_edge"], title_size=11, subtitle_size=6.5)
    draw_box(ax, 0.60, 0.105, 0.13, 0.085, "Pd nanoparticle\ndiameter", "Property occurrence", face=COLORS["property"], edge=COLORS["property_edge"], title_size=8.2, subtitle_size=6.7)
    draw_box(ax, 0.87, 0.105, 0.11, 0.085, "DataPoint", "synthetic result", face=COLORS["data"], edge=COLORS["data_edge"], title_size=8.8, subtitle_size=7.0)
    draw_box(ax, 0.705, 0.012, 0.125, 0.052, "M", "same measurement", face="#F8FAFB", edge=COLORS["measurement_edge"], title_size=7.8, subtitle_size=5.9, dashed=True)
    draw_box(ax, 0.855, 0.012, 0.125, 0.052, "A", "same analysis", face="#F8FAFB", edge=COLORS["process_edge"], title_size=7.8, subtitle_size=5.9, dashed=True)

    arrow(ax, (0.19, main_y), (0.215, 0.19), "usesInstrument", color=COLORS["measurement_edge"], fontsize=6.3, label_offset=(-0.035, 0))
    arrow(ax, (0.27, main_y), (0.375, 0.19), "hasParameter", color=COLORS["measurement_edge"], fontsize=6.3, label_offset=(0.025, 0))
    arrow(ax, (0.155, 0.305), (0.12, 0.19), "hasMetadata", color=COLORS["metadata_edge"], dashed=True, fontsize=6.3, label_offset=(-0.035, 0))
    arrow(ax, (0.942, main_y), (0.925, 0.19), "hasPart", color=COLORS["data_edge"], fontsize=6.3, label_offset=(0.035, 0))
    arrow(ax, (0.87, 0.148), (0.73, 0.148), "ofProperty", color=COLORS["data_edge"], fontsize=6.3, label_offset=(0, 0.025))
    arrow(ax, (0.60, 0.148), (0.57, 0.148), "hasQuantityValue", color=COLORS["property_edge"], fontsize=6.0, label_offset=(0, 0.026))
    arrow(ax, (0.89, 0.105), (0.79, 0.064), "fromMeasurement", color=COLORS["data_edge"], dashed=True, fontsize=5.8, label_offset=(-0.015, 0.018))
    arrow(ax, (0.95, 0.105), (0.93, 0.064), "prov:wasGeneratedBy", color=COLORS["data_edge"], dashed=True, fontsize=5.4, label_offset=(-0.045, 0.017))

    ax.text(0.02, 0.078, "Repeated M and A references denote the same occurrences shown above; they are not additional graph nodes.", fontsize=7.0, color=COLORS["muted"])
    save_figure(fig, "figure_5_synthetic_tem")


def compact_iri(value: str) -> str:
    replacements = {
        H2KG: "h2kg:",
        QUDT: "qudt:",
        "http://www.w3.org/ns/prov#": "prov:",
        "http://www.w3.org/2004/02/skos/core#": "skos:",
    }
    for prefix, compact in replacements.items():
        if value.startswith(prefix):
            return compact + value[len(prefix) :]
    return value


def markdown_table(headers: list[str], rows: Iterable[Iterable[object]]) -> str:
    def clean(value: object) -> str:
        return str(value if value is not None else "").replace("|", "\\|").replace("\n", " ")

    lines = ["| " + " | ".join(map(clean, headers)) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    lines.extend("| " + " | ".join(clean(value) for value in row) + " |" for row in rows)
    return "\n".join(lines)


def object_property_table() -> tuple[str, list[dict[str, str]]]:
    graph = Graph().parse(ONTOLOGY_RELEASE / "output/ontology/h2kg-v2-taxonomy.ttl", format="turtle")
    rows: list[dict[str, str]] = []
    for prop in graph.subjects(RDF.type, OWL.ObjectProperty):
        if not str(prop).startswith(H2KG):
            continue
        labels = [str(value) for value in graph.objects(prop, RDFS.label)]
        domains = [compact_iri(str(value)) for value in graph.objects(prop, RDFS.domain)]
        ranges = [compact_iri(str(value)) for value in graph.objects(prop, RDFS.range)]
        rows.append(
            {
                "property": compact_iri(str(prop)),
                "label": labels[0] if labels else str(prop).rsplit("#", 1)[-1],
                "domain": "; ".join(domains),
                "range": "; ".join(ranges),
            }
        )
    rows.sort(key=lambda row: row["property"].casefold())
    assert len(rows) == 23, f"Expected 23 object properties, found {len(rows)}"
    return markdown_table(["Property", "Label", "Domain", "Range"], ([row[k] for k in ("property", "label", "domain", "range")] for row in rows)), rows


def competency_table() -> str:
    report = json.loads((ONTOLOGY_RELEASE / "output/reports/taxonomy_validation.json").read_text(encoding="utf-8"))
    checks = report["competency_checks"]["checks"]
    rows = []
    for index, check in enumerate(checks, 1):
        rows.append((f"CQ{index}", check["id"].replace("-", " ").capitalize(), check["detail"], "Pass" if check["passed"] else "Fail", "; ".join(check.get("evidence", []))))
    assert len(rows) == 9 and all(row[3] == "Pass" for row in rows)
    return markdown_table(["ID", "Question/check", "Expected behavior", "Result", "Evidence"], rows)


def ontology_benchmark_table() -> tuple[str, list[dict[str, str]]]:
    path = ONTOLOGY_RELEASE / "output/benchmarks/tib_collections/ontology_benchmark.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 165, f"Expected 165 ontology-membership rows, found {len(rows)}"
    columns = [
        "Collection",
        "Ontology ID",
        "Ontology title",
        "Version",
        "TIB link",
        "TIB-reported terms",
        "H2KG classes assessed",
        "Exact IRI reuse",
        "Machine exact equivalent",
        "H2KG narrower",
        "H2KG broader",
        "Related non-equivalent",
        "Ambiguous abstained",
        "Unmatched",
    ]
    table = markdown_table(columns, ([row.get(column, "") for column in columns] for row in rows))
    return table, rows


def alignment_counts_table() -> tuple[str, list[dict[str, object]]]:
    payload = json.loads((ONTOLOGY_RELEASE / "output/ontology/external_alignment_registry.json").read_text(encoding="utf-8"))
    mappings = [row for row in payload["mappings"] if row.get("publication_status") == "asserted-machine-safe"]
    assert len(mappings) == 245
    grouped: dict[tuple[str, str, str], Counter[str]] = defaultdict(Counter)
    local_sets: dict[tuple[str, str, str], set[str]] = defaultdict(set)
    for row in mappings:
        key = (row.get("tib_ontology_id", ""), row.get("tib_ontology_title", ""), row.get("tib_ontology_version", ""))
        grouped[key][row["relation"]] += 1
        local_sets[key].add(row["local_iri"])
    output: list[dict[str, object]] = []
    for (ontology_id, title, version), counts in sorted(grouped.items(), key=lambda item: (item[0][1].casefold(), item[0][0])):
        output.append(
            {
                "ontology_id": ontology_id,
                "title": title,
                "version": version,
                "exact": counts["skos:exactMatch"],
                "broad": counts["skos:broadMatch"],
                "narrow": counts["skos:narrowMatch"],
                "total": sum(counts.values()),
                "h2kg_classes": len(local_sets[(ontology_id, title, version)]),
            }
        )
    assert len(output) == 29
    table = markdown_table(
        ["Ontology ID", "Ontology title", "Version", "Exact", "H₂KG narrower (`broadMatch`)", "H₂KG broader (`narrowMatch`)", "Assertions", "H₂KG classes"],
        ([row[key] for key in ("ontology_id", "title", "version", "exact", "broad", "narrow", "total", "h2kg_classes")] for row in output),
    )
    return table, output


def tem_field_table() -> tuple[str, list[dict[str, str]]]:
    path = ONTOLOGY_RELEASE / "output/method_profiles/tem/tem_field_dictionary.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    columns = ["h2kg_iri", "preferred_label", "semantic_role", "expected_value_form", "cardinality_or_condition", "qudt_guidance", "metadata_boundary"]
    display_headers = ["H₂KG IRI", "Preferred label", "Role", "Value form", "Cardinality/condition", "QUDT guidance", "Metadata boundary"]
    table = markdown_table(display_headers, ([compact_iri(row[column]) if column == "h2kg_iri" else row[column] for column in columns] for row in rows))
    return table, rows


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def build_supplement() -> None:
    object_table, property_rows = object_property_table()
    benchmark_table, benchmark_rows = ontology_benchmark_table()
    alignment_table, alignment_rows = alignment_counts_table()
    tem_table, tem_rows = tem_field_table()
    template = (HERE / "supplementary_information.template.md").read_text(encoding="utf-8")
    output = (
        template.replace("{{OBJECT_PROPERTIES_TABLE}}", object_table)
        .replace("{{COMPETENCY_TABLE}}", competency_table())
        .replace("{{ONTOLOGY_BENCHMARK_TABLE}}", benchmark_table)
        .replace("{{ALIGNMENT_COUNTS_TABLE}}", alignment_table)
        .replace("{{TEM_FIELD_TABLE}}", tem_table)
    )
    assert "{{" not in output
    (HERE / "supplementary_information.md").write_text(output, encoding="utf-8")
    write_csv(TABLES / "object_properties.csv", property_rows)
    write_csv(TABLES / "ontology_benchmark_165_rows.csv", benchmark_rows)
    write_csv(TABLES / "asserted_alignment_counts_by_ontology.csv", alignment_rows)
    write_csv(TABLES / "synthetic_tem_field_dictionary.csv", tem_rows)
    shutil.copyfile(
        ONTOLOGY_RELEASE / "output/benchmarks/tib_collections/h2kg_tib_manuscript_tables.xlsx",
        TABLES / "h2kg_tib_manuscript_tables.xlsx",
    )


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_inline(paragraph, text: str) -> None:
    token_re = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(45, 79, 91)
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def parse_table(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    def cells(line: str) -> list[str]:
        return [cell.strip().replace("\\|", "|") for cell in line.strip().strip("|").split("|")]

    header = cells(lines[0])
    rows = [cells(line) for line in lines[2:]]
    return header, rows


def configure_document(doc: Document, short_title: str) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    add_page_number(section)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [("Title", 20, "173042"), ("Heading 1", 15, "173042"), ("Heading 2", 12, "23689B"), ("Heading 3", 11, "337C8D")]:
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Normal" else "Aptos"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    if "Caption" in styles:
        styles["Caption"].font.name = "Aptos"
        styles["Caption"].font.size = Pt(9)
        styles["Caption"].font.color.rgb = RGBColor.from_string("4F6573")

    if "Code Block" not in [style.name for style in styles]:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code_style.font.name = "Consolas"
        code_style.font.size = Pt(8.5)
        code_style.paragraph_format.left_indent = Inches(0.25)
        code_style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.text = short_title
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("71828D")


def switch_orientation(doc: Document, landscape: bool) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section.left_margin = Inches(0.45)
        section.right_margin = Inches(0.45)
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
    add_page_number(section)


def markdown_to_docx(markdown_path: Path, output_path: Path, *, supplement: bool = False) -> None:
    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc, "H₂KG 2.2 supplementary information" if supplement else "H₂KG 2.2 domain ontology")
    index = 0
    in_code = False
    code_lines: list[str] = []
    benchmark_landscape = False
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if in_code:
                paragraph = doc.add_paragraph(style="Code Block")
                paragraph.add_run("\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        image_match = re.fullmatch(r"!\[(.*?)\]\((.*?)\)", line.strip())
        if image_match:
            image_path = (markdown_path.parent / image_match.group(2)).resolve()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(6.7 if not supplement else 6.4))
            index += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading = line[level:].strip()
            if supplement and heading.startswith("Supplementary Table S6") and not benchmark_landscape:
                switch_orientation(doc, True)
                benchmark_landscape = True
            elif supplement and benchmark_landscape and heading.startswith("Supplementary Table S7"):
                switch_orientation(doc, False)
                benchmark_landscape = False
            paragraph = doc.add_paragraph(style="Title" if level == 1 else f"Heading {min(level - 1, 3)}")
            add_inline(paragraph, heading)
            index += 1
            continue
        if line.startswith("| ") and index + 1 < len(lines) and re.match(r"^\|\s*:?-+", lines[index + 1]):
            table_lines = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            header, rows = parse_table(table_lines)
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for column, value in enumerate(header):
                cell = table.rows[0].cells[column]
                cell.text = value
                set_cell_shading(cell, "DCEAF7")
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(6.5 if len(header) > 8 else 8)
            for row_values in rows:
                cells = table.add_row().cells
                for column in range(len(header)):
                    value = row_values[column] if column < len(row_values) else ""
                    cells[column].text = value
                    cells[column].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    for paragraph in cells[column].paragraphs:
                        paragraph.paragraph_format.space_after = Pt(0)
                        for run in paragraph.runs:
                            run.font.size = Pt(5.7 if len(header) > 8 else 7.5)
                if len(table.rows) % 2 == 1:
                    for cell in cells:
                        set_cell_shading(cell, "F6F8F9")
            doc.add_paragraph()
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, line[2:].strip())
            index += 1
            continue
        paragraph_lines = [line.strip()]
        index += 1
        while index < len(lines) and lines[index].strip() and not lines[index].startswith(("#", "|", "- ", "```", "![")):
            paragraph_lines.append(lines[index].strip())
            index += 1
        text = " ".join(paragraph_lines)
        style = "Caption" if text.startswith("**Figure ") or text.startswith("**Table ") else None
        paragraph = doc.add_paragraph(style=style)
        if style == "Caption":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline(paragraph, text)

    core = doc.core_properties
    core.title = "Supplementary information for H₂KG 2.2" if supplement else "A domain-specific ontology for interoperable hydrogen electrochemical research workflows"
    core.subject = "H₂KG 2.2 ontology, validation and TIB benchmarking"
    core.author = "Marjan Kohandani; Mohammad J. Eslamibidgoli; Said Fathallah; Michael H. Eikerling; Kourosh Malek"
    core.keywords = "H₂KG, ontology, hydrogen electrochemistry, semantic interoperability, TIB Terminology Service"
    core.comments = "Generated reproducibly from the Markdown source in the AIMWORKS repository."
    doc.save(output_path)


def write_figure_sources() -> None:
    content = """# Figure content and source notes

All five figures were generated as editable SVG and 400 dpi PNG files by `build_article.py`. No third-party figure or unpublished source document is reproduced.

1. **Scope architecture:** generated from the H₂KG 2.2 release profile, taxonomy validation report, and external-alignment policy.
2. **Engineering workflow:** generated from the release pipeline stages and validation gates implemented under `ontology_release/src/aimworks_ontology_release/`.
3. **Taxonomy and role pattern:** branch counts come from `taxonomy_validation.json`; the occurrence pattern follows the canonical `hasParameter / ofProperty / hasQuantityValue` model.
4. **TIB benchmark:** percentages are read directly from the frozen `benchmark_summary.json` generated from the 14 October 2026 snapshot.
5. **Synthetic TEM workflow:** generated from `output/method_profiles/tem/synthetic_example.ttl` and the TEM SHACL profile. Every occurrence value is synthetic.

The SVG files retain text and shapes as editable vector objects. Arrows in Figure 5 follow RDF subject-to-object direction. The edge audit is stored in `tables/figure_5_edge_audit.csv`.
"""
    (HERE / "figure_content_and_sources.md").write_text(content, encoding="utf-8")
    edges = [
        {"subject": "Mix + sonicate", "predicate": "hasInputMaterial", "object": "Pd/C"},
        {"subject": "Mix + sonicate", "predicate": "hasInputMaterial", "object": "Liquid mixture"},
        {"subject": "Mix + sonicate", "predicate": "hasOutputMaterial", "object": "Dispersion"},
        {"subject": "Deposit + dry", "predicate": "hasInputMaterial", "object": "Dispersion"},
        {"subject": "Deposit + dry", "predicate": "hasInputMaterial", "object": "TEM grid"},
        {"subject": "Deposit + dry", "predicate": "hasOutputMaterial", "object": "TEM sample"},
        {"subject": "TEM measurement", "predicate": "hasInputMaterial", "object": "TEM sample"},
        {"subject": "TEM measurement", "predicate": "usesInstrument", "object": "TEM instrument"},
        {"subject": "TEM measurement", "predicate": "hasParameter", "object": "Parameter records"},
        {"subject": "TEM measurement", "predicate": "hasMetadata", "object": "Metadata"},
        {"subject": "TEM measurement", "predicate": "hasOutputData", "object": "Raw TEM images"},
        {"subject": "Format conversion", "predicate": "hasInputData", "object": "Raw TEM images"},
        {"subject": "Format conversion", "predicate": "hasOutputData", "object": "Processed images"},
        {"subject": "Particle analysis", "predicate": "hasInputData", "object": "Processed images"},
        {"subject": "Particle analysis", "predicate": "hasOutputData", "object": "Result data"},
        {"subject": "Result data", "predicate": "hasPart", "object": "DataPoint"},
        {"subject": "DataPoint", "predicate": "ofProperty", "object": "Pd nanoparticle diameter occurrence"},
        {"subject": "Pd nanoparticle diameter occurrence", "predicate": "hasQuantityValue", "object": "Synthetic 5 nm QUDT value"},
        {"subject": "DataPoint", "predicate": "fromMeasurement", "object": "TEM measurement"},
        {"subject": "DataPoint", "predicate": "prov:wasGeneratedBy", "object": "Particle analysis"},
    ]
    write_csv(TABLES / "figure_5_edge_audit.csv", edges)


def bibliography_checks() -> dict[str, object]:
    manuscript = (HERE / "manuscript.md").read_text(encoding="utf-8")
    body, reference_section = manuscript.split("## References", 1)
    references = [
        {"number": int(match.group(1)), "citation": match.group(2).strip()}
        for match in re.finditer(r"(?m)^(\d+)\.\s+(.+)$", reference_section)
    ]
    cited: set[int] = set()
    for group in re.findall(r"\[([0-9,–\- ]+)\]", body):
        for item in re.split(r"[, ]+", group):
            if not item:
                continue
            if "–" in item or "-" in item:
                start, end = (int(value) for value in re.split(r"[–-]", item))
                cited.update(range(start, end + 1))
            else:
                cited.add(int(item))
    numbers = [row["number"] for row in references]
    links = {
        row["number"]: re.findall(r"https://[^\s)]+", row["citation"])
        for row in references
    }
    return {
        "references": references,
        "sequential": numbers == list(range(1, len(numbers) + 1)),
        "duplicate_numbers": sorted(number for number, count in Counter(numbers).items() if count > 1),
        "uncited_numbers": sorted(set(numbers) - cited),
        "missing_reference_numbers": sorted(cited - set(numbers)),
        "references_without_persistent_link": sorted(number for number, urls in links.items() if not urls),
        "links": links,
    }


def write_bibliography_verification() -> None:
    checks = bibliography_checks()
    rows = []
    for reference in checks["references"]:
        number = reference["number"]
        citation = reference["citation"]
        link = checks["links"][number][0]
        citation_without_link = citation.replace(link, "").strip()
        rows.append((number, citation_without_link, link, "Cited; persistent link present"))
    content = """# Bibliography verification

The article bibliography was checked on 19 October 2026 against the cited DOI, standards-body, institutional-repository, or project documentation records. The reproducible builder additionally verifies that numbering is consecutive, every entry is cited in the manuscript, every in-text citation has a bibliography entry, and every entry contains a persistent HTTPS link.

""" + markdown_table(["No.", "Bibliographic record", "Persistent record", "Internal check"], rows) + "\n"
    (HERE / "bibliography_verification.md").write_text(content, encoding="utf-8")


def validate() -> dict[str, object]:
    manuscript = (HERE / "manuscript.md").read_text(encoding="utf-8")
    supplement = (HERE / "supplementary_information.md").read_text(encoding="utf-8")
    title = manuscript.splitlines()[0].lstrip("# ").strip()
    abstract = manuscript.split("## Abstract", 1)[1].split("## Introduction", 1)[0]
    abstract_words = re.findall(r"\b[\w₂–-]+\b", abstract)
    heading_positions = [manuscript.index(f"## {heading}") for heading in ("Introduction", "Methods", "Results", "Discussion")]

    banned = ["EMMO-aligned", "FAIR-compliant", "domain-expert-reviewed", "AI-reviewed", "expert-draft", "release candidate", "source_grounded", "source-grounded"]
    banned_hits = [term for term in banned if term.casefold() in (manuscript + supplement).casefold()]
    scientific_body = manuscript.split("## Acknowledgements", 1)[0]
    decode_scientific_hits = len(re.findall(r"\bDECODE\b", scientific_body, flags=re.IGNORECASE))

    benchmark = json.loads((ONTOLOGY_RELEASE / "output/benchmarks/tib_collections/benchmark_summary.json").read_text(encoding="utf-8"))
    collection_totals = {
        row["collection"]: sum(row[f"{name}_count"] for name in ("exact_iri_reuse", "machine_exact_equivalent", "h2kg_narrower", "h2kg_broader", "related_non_equivalent", "ambiguous_abstained", "unmatched"))
        for row in benchmark["collection_summary"]
    }
    figure_files = [FIGURES / f"figure_{index}_{stem}.{suffix}" for index, stem in [(1, "scope_architecture"), (2, "engineering_workflow"), (3, "taxonomy_role_pattern"), (4, "tib_benchmark"), (5, "synthetic_tem")] for suffix in ("svg", "png")]

    docx_paths = [HERE / "H2KG_2_2_Scientific_Data_Article.docx", HERE / "H2KG_2_2_Supplementary_Information.docx"]
    docx_valid = {}
    for path in docx_paths:
        with zipfile.ZipFile(path) as archive:
            corrupt = archive.testzip()
        Document(path)
        docx_valid[path.name] = corrupt is None

    tem_text = " ".join(
        path.read_text(encoding="utf-8", errors="ignore")
        for path in [
            ONTOLOGY_RELEASE / "output/examples/tem_pilot/tem_pilot_example.jsonld",
            ONTOLOGY_RELEASE / "output/method_profiles/tem/synthetic_example.ttl",
            ONTOLOGY_RELEASE / "output/method_profiles/tem/validation_report.json",
        ]
    ).casefold()
    tem_forbidden = [token for token in ("doi.org/10.3390", "@gmail", "@uconn", "source_grounded", "source-grounded") if token in tem_text]
    bibliography = bibliography_checks()

    checks = {
        "title_characters": len(title),
        "title_under_110": len(title) < 110,
        "abstract_words": len(abstract_words),
        "abstract_at_most_170": len(abstract_words) <= 170,
        "imrad_order": heading_positions == sorted(heading_positions),
        "banned_wording_hits": banned_hits,
        "decode_scientific_content_hits": decode_scientific_hits,
        "benchmark_deprecated_count": benchmark["deprecated_h2kg_classes_excluded"],
        "benchmark_collection_totals": collection_totals,
        "benchmark_totals_valid": all(value == 2221 for value in collection_totals.values()),
        "ontology_rows": benchmark["ontology_rows"],
        "all_figures_present": all(path.exists() and path.stat().st_size > 1000 for path in figure_files),
        "docx_valid": docx_valid,
        "tem_forbidden_evidence_hits": tem_forbidden,
        "synthetic_tem_label_count": len(re.findall(r"\bsynthetic\b", manuscript, flags=re.IGNORECASE)),
        "comment_resolution_rows": sum(1 for line in (HERE / "review_comment_resolution.md").read_text(encoding="utf-8").splitlines() if re.match(r"^\| \d+ ", line)),
        "bibliography_references": len(bibliography["references"]),
        "bibliography_sequential": bibliography["sequential"],
        "bibliography_duplicate_numbers": bibliography["duplicate_numbers"],
        "bibliography_uncited_numbers": bibliography["uncited_numbers"],
        "bibliography_missing_reference_numbers": bibliography["missing_reference_numbers"],
        "bibliography_references_without_persistent_link": bibliography["references_without_persistent_link"],
    }
    failures = []
    if not checks["title_under_110"]:
        failures.append("title length")
    if not checks["abstract_at_most_170"]:
        failures.append("abstract length")
    if not checks["imrad_order"]:
        failures.append("IMRaD order")
    if banned_hits:
        failures.append("banned wording")
    if decode_scientific_hits:
        failures.append("DECODE scientific content")
    if checks["benchmark_deprecated_count"] != 165 or checks["ontology_rows"] != 165 or not checks["benchmark_totals_valid"]:
        failures.append("benchmark integrity")
    if not checks["all_figures_present"] or not all(docx_valid.values()):
        failures.append("deliverable integrity")
    if tem_forbidden:
        failures.append("TEM sanitization")
    if checks["comment_resolution_rows"] != 31:
        failures.append("comment-resolution completeness")
    if not bibliography["sequential"] or bibliography["duplicate_numbers"] or bibliography["uncited_numbers"] or bibliography["missing_reference_numbers"] or bibliography["references_without_persistent_link"]:
        failures.append("bibliography integrity")
    checks["conforms"] = not failures
    checks["failures"] = failures
    (HERE / "manuscript_validation.json").write_text(json.dumps(checks, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    if failures:
        raise RuntimeError("Manuscript validation failed: " + ", ".join(failures))
    return checks


def build() -> None:
    configure_plotting()
    FIGURES.mkdir(parents=True, exist_ok=True)
    TABLES.mkdir(parents=True, exist_ok=True)
    figure_1_scope_architecture()
    figure_2_engineering_workflow()
    figure_3_taxonomy_role_pattern()
    figure_4_tib_benchmark()
    figure_5_synthetic_tem()
    build_supplement()
    write_figure_sources()
    write_bibliography_verification()
    markdown_to_docx(HERE / "manuscript.md", HERE / "H2KG_2_2_Scientific_Data_Article.docx")
    markdown_to_docx(HERE / "supplementary_information.md", HERE / "H2KG_2_2_Supplementary_Information.docx", supplement=True)
    report = validate()
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    build()
